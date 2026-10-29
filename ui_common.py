"""
Shared look and controls for the dashboard's secondary pages (Ticket Trend
Analysis, Resolved Tickets History): palette, base CSS, page header, the
Day / Week (Sun-Sat) / Custom period picker, the Jira sync trigger and the
Excel styling.
"""

import threading
from datetime import datetime, timedelta

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill

import config
import trend_analysis as ta
import ui_table
import trend_store

SURFACE_PAGE = "#0d0d14"
SURFACE_CARD = "#14141c"
SURFACE_PANEL = "#1a1a24"
HEADER_BG = "#20202c"
TEXT_PRIMARY = "#F5F5F7"
TEXT_SECONDARY = "#C3C2B7"
TEXT_MUTED = "#898781"
ACCENT_FROM = "#7C3AED"
ACCENT_TO = "#DB2777"
COLOR_GOOD = "#0ca30c"
COLOR_WARNING = "#fab219"
COLOR_CRITICAL = "#d03b3b"

BASE_CSS = f"""
    #MainMenu, footer, header {{visibility: hidden;}}
    [data-testid="stHeader"], [data-testid="stSidebar"], [data-testid="stSidebarCollapsedControl"] {{ display: none; }}
    [data-testid="stElementContainer"]:has(iframe[title="streamlit_autorefresh.st_autorefresh"]),
    [data-testid="stElementContainer"]:has(style) {{ display: none; }}
    [data-testid="stAppViewContainer"] {{
        background:
            radial-gradient(circle at 0% 0%, rgba(124,58,237,0.22), transparent 45%),
            radial-gradient(circle at 100% 0%, rgba(31,111,235,0.18), transparent 45%),
            {SURFACE_PAGE};
    }}
    [data-testid="stMain"] {{ background: transparent; overflow-x: hidden; }}
    [data-stale="true"] {{ opacity: 1 !important; transition: none !important; }}
    .block-container, [data-testid="stMainBlockContainer"] {{
        background: {SURFACE_CARD}; border: 1px solid rgba(255,255,255,0.07); border-radius: 20px;
        padding: 0.9rem 1.8rem 1.6rem 1.8rem !important; margin-top: 0.6rem; margin-bottom: 0.8rem;
        max-width: min(1880px, calc(100vw - 2.5rem)); box-shadow: 0 20px 60px rgba(0,0,0,0.45);
    }}
    .app-header {{ display: flex; align-items: center; gap: 10px; }}
    .app-logo {{
        width: 36px; height: 36px; border-radius: 10px; flex-shrink: 0;
        background: linear-gradient(135deg, {ACCENT_FROM}, {ACCENT_TO});
        display: flex; align-items: center; justify-content: center;
    }}
    .app-title {{ font-size: 1.15rem; font-weight: 700; color: {TEXT_PRIMARY}; line-height: 1.1; }}
    .app-subtitle {{ font-size: 0.75rem; color: {TEXT_MUTED}; margin-top: 1px; }}
    .sync-note {{ font-size: 0.68rem; color: {TEXT_MUTED}; text-align: right; }}
    .section-title {{
        font-size: 0.72rem; font-weight: 700; letter-spacing: 0.06em; text-transform: uppercase;
        color: {TEXT_SECONDARY}; background: rgba(255,255,255,0.04);
        border: 1px solid rgba(255,255,255,0.08); border-radius: 8px; padding: 7px 14px; margin: 6px 0 2px 0;
    }}
    .kpi {{
        border: 1px solid rgba(255,255,255,0.08); border-radius: 12px; background: {SURFACE_PANEL};
        padding: 10px 14px; border-top: 3px solid var(--cat);
    }}
    .kpi-label {{ font-size: 0.62rem; font-weight: 700; letter-spacing: 0.05em; text-transform: uppercase; color: {TEXT_MUTED}; }}
    .kpi-value {{ font-size: 1.7rem; font-weight: 700; color: {TEXT_PRIMARY}; line-height: 1.15; }}
    .kpi-sub {{ font-size: 0.72rem; color: {TEXT_SECONDARY}; }}
    .delta-bad {{ color: {COLOR_CRITICAL}; font-weight: 700; }}
    .delta-good {{ color: {COLOR_GOOD}; font-weight: 700; }}
    .delta-flat {{ color: {TEXT_MUTED}; font-weight: 700; }}
    button[data-baseweb="tab"] {{
        background-color: transparent; border-radius: 999px !important; padding: 6px 20px !important;
        color: {TEXT_MUTED} !important; margin-right: 4px;
    }}
    button[data-baseweb="tab"][aria-selected="true"] {{
        background: linear-gradient(135deg, {ACCENT_FROM}, {ACCENT_TO}) !important; color: white !important; font-weight: 600;
    }}
    div[data-baseweb="tab-list"] {{
        gap: 2px; border-bottom: none; background: rgba(255,255,255,0.04);
        border-radius: 999px; padding: 6px; width: fit-content;
    }}
    div[data-baseweb="tab-highlight"], div[data-baseweb="tab-border"] {{ display: none; }}
    div[data-testid="stDataFrame"] {{ border: 1px solid rgba(255,255,255,0.07); border-radius: 10px; overflow: hidden; }}
    .app-loader-status {{
        display: flex; align-items: center; gap: 10px; margin: 26px 0 18px 0;
        font-size: 0.9rem; color: {TEXT_SECONDARY};
    }}
    .app-loader-spinner {{
        width: 18px; height: 18px; border-radius: 50%; display: inline-block; flex-shrink: 0;
        border: 2.5px solid rgba(255,255,255,0.15); border-top-color: {ACCENT_TO};
        animation: app-loader-spin 0.9s linear infinite;
    }}
    @keyframes app-loader-spin {{ to {{ transform: rotate(360deg); }} }}
"""


def apply_css(extra_css=""):
    st.markdown(f"<style>{BASE_CSS}{ui_table.TABLE_CSS}{extra_css}</style>", unsafe_allow_html=True)


def page_header(title, subtitle, icon_svg_paths):
    """Logo + title on the left; returns the right-hand column (sync button / note)."""
    left, right = st.columns([4, 1.3])
    left.markdown(
        '<div class="app-header"><div class="app-logo">'
        '<svg viewBox="0 0 24 24" width="19" height="19" fill="none" stroke="white" stroke-width="2.2" '
        f'stroke-linecap="round" stroke-linejoin="round">{icon_svg_paths}</svg></div>'
        f'<div><div class="app-title">{title}</div><div class="app-subtitle">{subtitle}</div></div></div>',
        unsafe_allow_html=True,
    )
    return right


def loading(text):
    """A "spinner + text" status line (same look as the dashboard's loading
    screen) in a placeholder; call .empty() on the result once done."""
    slot = st.empty()
    slot.markdown(f'<div class="app-loader-status"><span class="app-loader-spinner"></span>{text}</div>',
                  unsafe_allow_html=True)
    return slot


def loading_if_new(state_key, params, text):
    """A placeholder in which ui.loading's status is shown only when params
    differ from this session's last run of the same step (real work is about
    to happen), so cached reruns like the 5-minute auto-refresh don't flash it.
    The slot is created on every run and cleared by done_loading(): a run that
    Streamlit cuts short (e.g. the auto-refresh's first rerun) leaves its status
    behind, and the next run's slot replaces it."""
    slot = st.empty()
    if st.session_state.get(state_key) != params:
        st.session_state[state_key] = params
        slot.markdown(f'<div class="app-loader-status"><span class="app-loader-spinner"></span>{text}</div>',
                      unsafe_allow_html=True)
    return slot


def done_loading(slot):
    slot.empty()


def jira_url(key):
    return f'{config.JIRA_BASE_URL.rstrip("/")}/browse/{key}'


# ---------------------------------------------------------------------------
# Data freshness: the local cache is kept fresh by `python trend_store.py
# --loop`; if that agent isn't running, a page syncs itself once it's stale.
# trend_store.sync() allows one sync across all processes; this in-process
# lock just stops sessions of one server queueing up behind each other.
# ---------------------------------------------------------------------------
@st.cache_resource
def _sync_lock():
    return threading.Lock()


def ensure_fresh(force=False):
    last = trend_store.last_sync_time()
    stale = last is None or datetime.now() - last > timedelta(minutes=config.TREND_STALE_MINUTES)
    if not (force or stale):
        return None
    lock = _sync_lock()
    if trend_store.sync_running() or not lock.acquire(blocking=False):
        if force:
            st.info("A Jira sync is already running (another page or the sync agent); "
                    "the data will update when it finishes.")
        return None
    try:
        if last is None or trend_store.needs_full_reload():
            bar = st.progress(0.0, text=f"Loading the last {ta.HISTORY_LABEL} of tickets from Jira (about 10 minutes)…")
            msg = trend_store.sync(lambda f, t: bar.progress(min(f, 1.0), text=t))
            bar.empty()
        else:
            msg = trend_store.sync()
        return msg
    except Exception as exc:
        st.warning(f"Jira sync failed, showing cached data: {exc}")
        return None
    finally:
        lock.release()


# ---------------------------------------------------------------------------
# Period picker
# ---------------------------------------------------------------------------
DAY, WEEK, CUSTOM = "Day", "Week (Sun – Sat)", "Custom range"


def period_picker(today, key, widths=(1.6, 1.6, 1.2)):
    """Day (default: today) / Sunday-Saturday week / custom range, plus the
    comparison period. Returns a dict of mode, start, end, label, prev_range,
    prev_label; the caller gets one spare column via "extra_col"."""
    earliest = today - timedelta(days=config.TREND_HISTORY_DAYS)
    c_mode, c_pick, c_extra = st.columns(list(widths))
    with c_mode:
        mode = st.segmented_control("View", [DAY, WEEK, CUSTOM], default=DAY, key=f"{key}_mode") or DAY
    with c_pick:
        if mode == DAY:
            day = st.date_input("Date", value=today, min_value=earliest, max_value=today,
                                format="DD/MM/YYYY", key=f"{key}_day")
            start = end = day
            label = f"{day:%d %b %Y}" + (" (today)" if day == today else "")
            prev_range = ta.previous_period(start, end)
            prev_label = "previous day"
        elif mode == WEEK:
            weeks_back = [ta.week_start(today) - timedelta(weeks=i) for i in range(config.TREND_HISTORY_DAYS // 7)]
            start = st.selectbox("Week", weeks_back, key=f"{key}_week", format_func=lambda d: ta.week_label(d)
                                 + (" (this week)" if d == ta.week_start(today) else ""))
            end = min(start + timedelta(days=6), today)
            label = ta.week_label(start)
            prev_range = ta.previous_period(start, end, week=True)
            prev_label = ("previous week" if (end - start).days == 6
                          else f"same days last week ({prev_range[0]:%d %b} – {prev_range[1]:%d %b})")
        else:
            picked = st.date_input("Date range", value=(today - timedelta(days=6), today), key=f"{key}_range",
                                   min_value=earliest, max_value=today, format="DD/MM/YYYY")
            start, end = (picked[0], picked[-1]) if isinstance(picked, (tuple, list)) and picked else (today, today)
            label = f"{start:%d %b %Y} – {end:%d %b %Y}"
            prev_range = ta.previous_period(start, end)
            prev_label = f"previous {(end - start).days + 1} days"
    return {"mode": mode, "start": start, "end": end, "label": label, "prev_range": prev_range,
            "prev_label": prev_label, "extra_col": c_extra}


def delta_html(cur, prev, higher_is_good=False):
    if prev == 0:
        return '<span class="delta-flat">new</span>' if cur else '<span class="delta-flat">—</span>'
    pct = (cur - prev) / prev * 100
    good = pct > 0 if higher_is_good else pct < 0
    cls = "delta-flat" if pct == 0 else "delta-good" if good else "delta-bad"
    arrow = "▲" if pct > 0 else "▼" if pct < 0 else "■"
    return f'<span class="{cls}">{arrow} {abs(pct):.0f}%</span>'


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
LINK_COLUMNS = ("Link",)                # cells holding a URL become clickable
LINKED_ID_COLUMNS = ("Key", "Ticket")   # ... and so does the ticket key beside it


def write_excel(sheets):
    """{sheet name: DataFrame} -> styled .xlsx bytes (dark header, wrapped
    cells, sensible widths, frozen header row). A "Link" column and the
    "Key" / "Ticket" column in the same row become clickable Jira links."""
    import io
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            header_fill = PatternFill("solid", fgColor="20202C")
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
            for idx, col_name in enumerate(frame.columns, start=1):
                longest = frame[col_name].astype(str).str.len().max() if len(frame) else 0
                width = min(70, max(len(str(col_name)), longest or 0) + 2)
                ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            cols = list(frame.columns)
            # SLA % in the dashboard's colors: >= 90 green, >= 70 amber, else red.
            for c in cols:
                if str(c).strip().lower().endswith("sla %"):
                    for r, v in enumerate(frame[c], start=2):
                        if isinstance(v, (int, float)) and not pd.isna(v):
                            color = "0CA30C" if v >= 90 else "C98500" if v >= 70 else "D03B3B"
                            cell = ws.cell(row=r, column=cols.index(c) + 1)
                            cell.font = Font(bold=True, color=color)
                            cell.number_format = '0.0"%"'
            link_col = next((c for c in LINK_COLUMNS if c in cols), None)
            if link_col:
                targets = [cols.index(link_col) + 1] + [cols.index(c) + 1 for c in LINKED_ID_COLUMNS if c in cols]
                for r, url in enumerate(frame[link_col], start=2):
                    if isinstance(url, str) and url.startswith("http"):
                        for c in targets:
                            cell = ws.cell(row=r, column=c)
                            cell.hyperlink = url
                            cell.font = Font(color="0563C1", underline="single")
            ws.freeze_panes = "A2"
    return buf.getvalue()
