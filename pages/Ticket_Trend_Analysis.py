"""
Ticket Trend Analysis: top issue drivers for SAC (Retail + SAC), B2B and
Order Fallout over a day, a Sunday-Saturday week or a custom range, weekly
trends, and the SAL / SFS fixes behind each driver.

Opened in a new browser tab from the dashboard's "Ticket Trend Analysis"
button. Data comes from the local cache kept by trend_store.py.
"""

import html
import io
import threading
from datetime import date, datetime, timedelta

import altair as alt
import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from streamlit_autorefresh import st_autorefresh

import config
import trend_analysis as ta
import trend_store

SURFACE_PAGE = "#0d0d14"
SURFACE_CARD = "#14141c"
SURFACE_PANEL = "#1a1a24"
HEADER_BG = "#20202c"
TEXT_PRIMARY = "#F5F5F7"
TEXT_SECONDARY = "#C3C2B7"
TEXT_MUTED = "#898781"
ACCENT_FROM = "#7C3AED"
ACCENT_TO = "#DB2777"
COLOR_GOOD = "#0ca30c"
COLOR_WARNING = "#fab219"
COLOR_CRITICAL = "#d03b3b"
CATEGORY_COLOR = {"SAC": "#d95926", "B2B": "#199e70", "Order Fallout": "#c98500"}
# Fixed categorical order (dark-surface steps) for the top-5 driver lines.
SERIES_COLORS = ["#3987e5", "#d95926", "#199e70", "#c98500", "#d55181"]

FIX_STATE_COLOR = {
    "deployed": COLOR_GOOD, "resolved": COLOR_GOOD, "awaiting_release": "#3987e5",
    "in_progress": COLOR_WARNING, "not_started": TEXT_MUTED, "open": COLOR_WARNING,
    "cancelled": "#5b5b56",
}

st.set_page_config(page_title="Ticket Trend Analysis", page_icon="📈", layout="wide")
st_autorefresh(interval=5 * 60 * 1000, key="trend_autorefresh")

st.markdown(
    f"""
    <style>
        #MainMenu, footer, header {{visibility: hidden;}}
        [data-testid="stHeader"], [data-testid="stSidebar"], [data-testid="stSidebarCollapsedControl"] {{ display: none; }}
        [data-testid="stElementContainer"]:has(iframe[title="streamlit_autorefresh.st_autorefresh"]),
        [data-testid="stElementContainer"]:has(style) {{ display: none; }}
        [data-testid="stAppViewContainer"] {{
            background:
                radial-gradient(circle at 0% 0%, rgba(124,58,237,0.22), transparent 45%),
                radial-gradient(circle at 100% 0%, rgba(31,111,235,0.18), transparent 45%),
                {SURFACE_PAGE};
        }}
        [data-testid="stMain"] {{ background: transparent; overflow-x: hidden; }}
        [data-stale="true"] {{ opacity: 1 !important; transition: none !important; }}
        .block-container, [data-testid="stMainBlockContainer"] {{
            background: {SURFACE_CARD}; border: 1px solid rgba(255,255,255,0.07); border-radius: 20px;
            padding: 0.9rem 1.8rem 1.6rem 1.8rem !important; margin-top: 0.6rem; margin-bottom: 0.8rem;
            max-width: min(1880px, calc(100vw - 2.5rem)); box-shadow: 0 20px 60px rgba(0,0,0,0.45);
        }}
        .app-header {{ display: flex; align-items: center; gap: 10px; }}
        .app-logo {{
            width: 36px; height: 36px; border-radius: 10px; flex-shrink: 0;
            background: linear-gradient(135deg, {ACCENT_FROM}, {ACCENT_TO});
            display: flex; align-items: center; justify-content: center;
        }}
        .app-title {{ font-size: 1.15rem; font-weight: 700; color: {TEXT_PRIMARY}; line-height: 1.1; }}
        .app-subtitle {{ font-size: 0.75rem; color: {TEXT_MUTED}; margin-top: 1px; }}
        .sync-note {{ font-size: 0.68rem; color: {TEXT_MUTED}; text-align: right; }}
        .section-title {{
            font-size: 0.72rem; font-weight: 700; letter-spacing: 0.06em; text-transform: uppercase;
            color: {TEXT_SECONDARY}; background: rgba(255,255,255,0.04);
            border: 1px solid rgba(255,255,255,0.08); border-radius: 8px; padding: 7px 14px; margin: 6px 0 2px 0;
        }}
        .kpi {{
            border: 1px solid rgba(255,255,255,0.08); border-radius: 12px; background: {SURFACE_PANEL};
            padding: 10px 14px; border-top: 3px solid var(--cat);
        }}
        .kpi-label {{ font-size: 0.62rem; font-weight: 700; letter-spacing: 0.05em; text-transform: uppercase; color: {TEXT_MUTED}; }}
        .kpi-value {{ font-size: 1.7rem; font-weight: 700; color: {TEXT_PRIMARY}; line-height: 1.15; }}
        .kpi-sub {{ font-size: 0.72rem; color: {TEXT_SECONDARY}; }}
        .delta-up {{ color: {COLOR_CRITICAL}; font-weight: 700; }}
        .delta-down {{ color: {COLOR_GOOD}; font-weight: 700; }}
        .delta-flat {{ color: {TEXT_MUTED}; font-weight: 700; }}
        .alert-box {{
            border-radius: 10px; padding: 8px 14px; margin: 4px 0; font-size: 0.82rem; color: {TEXT_PRIMARY};
            border: 1px solid; display: flex; gap: 10px; align-items: baseline;
        }}
        .alert-recurring {{ background: rgba(208,59,59,0.12); border-color: rgba(208,59,59,0.45); }}
        .alert-overdue {{ background: rgba(250,178,25,0.10); border-color: rgba(250,178,25,0.4); }}
        .alert-tag {{ font-size: 0.62rem; font-weight: 700; text-transform: uppercase; letter-spacing: 0.05em;
                      padding: 2px 8px; border-radius: 999px; white-space: nowrap; }}
        .alert-recurring .alert-tag {{ background: {COLOR_CRITICAL}; color: white; }}
        .alert-overdue .alert-tag {{ background: {COLOR_WARNING}; color: #1a1a1a; }}
        .driver-card {{
            border: 1px solid rgba(255,255,255,0.08); border-radius: 12px; background: {SURFACE_PANEL};
            padding: 12px 16px; margin-bottom: 4px;
        }}
        .driver-head {{ display: flex; align-items: baseline; gap: 12px; }}
        .driver-rank {{ font-size: 0.8rem; font-weight: 700; color: {TEXT_MUTED}; }}
        .driver-name {{ font-size: 0.95rem; font-weight: 700; color: {TEXT_PRIMARY}; flex: 1; word-break: break-word; }}
        .driver-count {{ font-size: 1.15rem; font-weight: 700; color: {TEXT_PRIMARY}; white-space: nowrap; }}
        .driver-meta {{ font-size: 0.72rem; color: {TEXT_MUTED}; margin-top: 2px; }}
        .driver-summary {{ font-size: 0.8rem; color: {TEXT_SECONDARY}; margin: 8px 0 6px 0; line-height: 1.45; }}
        .fix-row {{ display: flex; flex-wrap: wrap; align-items: center; gap: 8px; font-size: 0.78rem;
                    color: {TEXT_SECONDARY}; padding: 5px 0; border-top: 1px solid rgba(255,255,255,0.06); }}
        .fix-proj {{ font-size: 0.62rem; font-weight: 700; color: {TEXT_MUTED}; width: 30px; }}
        .fix-key {{ color: #6ea8ff; font-weight: 700; text-decoration: none; }}
        .fix-key:hover {{ text-decoration: underline; }}
        .fix-chip {{ font-size: 0.66rem; font-weight: 700; padding: 1px 8px; border-radius: 999px; color: #111; white-space: nowrap; }}
        .fix-match {{ font-size: 0.66rem; color: {TEXT_MUTED}; border: 1px solid rgba(255,255,255,0.15);
                      padding: 0 6px; border-radius: 999px; }}
        .fix-summary {{ color: {TEXT_MUTED}; flex: 1; min-width: 200px; overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }}
        .fix-none {{ color: {TEXT_MUTED}; font-style: italic; }}
        .fix-alert {{ width: 100%; font-weight: 600; }}
        .fix-alert.recurring {{ color: #ff8a8a; }}
        .fix-alert.overdue {{ color: {COLOR_WARNING}; }}
        button[data-baseweb="tab"] {{
            background-color: transparent; border-radius: 999px !important; padding: 6px 20px !important;
            color: {TEXT_MUTED} !important; margin-right: 4px;
        }}
        button[data-baseweb="tab"][aria-selected="true"] {{
            background: linear-gradient(135deg, {ACCENT_FROM}, {ACCENT_TO}) !important; color: white !important; font-weight: 600;
        }}
        div[data-baseweb="tab-list"] {{
            gap: 2px; border-bottom: none; background: rgba(255,255,255,0.04);
            border-radius: 999px; padding: 6px; width: fit-content;
        }}
        div[data-baseweb="tab-highlight"], div[data-baseweb="tab-border"] {{ display: none; }}
        div[data-testid="stDataFrame"] {{ border: 1px solid rgba(255,255,255,0.07); border-radius: 10px; overflow: hidden; }}
    </style>
    """,
    unsafe_allow_html=True,
)


# ---------------------------------------------------------------------------
# Data: local cache, kept fresh by `python trend_store.py --loop`; if that
# agent isn't running, the page syncs itself once the cache is stale.
# ---------------------------------------------------------------------------
@st.cache_resource
def _sync_lock():
    return threading.Lock()


def ensure_fresh(force=False):
    last = trend_store.last_sync_time()
    stale = last is None or datetime.now() - last > timedelta(minutes=config.TREND_STALE_MINUTES)
    if not (force or stale):
        return None
    lock = _sync_lock()
    if not lock.acquire(blocking=False):
        return None  # another session is syncing
    try:
        if last is None:
            bar = st.progress(0.0, text="First run: loading the last 6 months of tickets from Jira (a few minutes)…")
            msg = trend_store.sync(lambda f, t: bar.progress(min(f, 1.0), text=t))
            bar.empty()
        else:
            msg = trend_store.sync()
        return msg
    except Exception as exc:
        st.warning(f"Jira sync failed, showing cached data: {exc}")
        return None
    finally:
        lock.release()


@st.cache_data(show_spinner="Analysing tickets…")
def load_data(sync_stamp):
    return ta.load_tickets(), ta.load_fixes()


def jira_url(key):
    return f'{config.JIRA_BASE_URL.rstrip("/")}/browse/{key}'


# ---------------------------------------------------------------------------
# Header
# ---------------------------------------------------------------------------
head_l, head_r = st.columns([4, 1.3])
with head_l:
    st.markdown(
        '<div class="app-header"><div class="app-logo">'
        '<svg viewBox="0 0 24 24" width="19" height="19" fill="none" stroke="white" stroke-width="2.2" '
        'stroke-linecap="round" stroke-linejoin="round"><path d="M3 17l6-6 4 4 8-8"/><path d="M15 7h6v6"/></svg>'
        '</div><div><div class="app-title">Ticket Trend Analysis</div>'
        '<div class="app-subtitle">Top issue drivers for SAC (Retail + SAC), B2B and Order Fallout, '
        'with SAL / SFS fix tracking</div></div></div>',
        unsafe_allow_html=True,
    )
with head_r:
    force = st.button("🔄  Sync with Jira now", width="stretch")

ensure_fresh(force=force)
last_sync = trend_store.last_sync_time()
if last_sync is None:
    st.error("No ticket data yet: the Jira sync has not completed. Check the Jira credentials in .env and retry.")
    st.stop()

tickets, fixes = load_data(last_sync.isoformat())
with head_r:
    st.markdown(
        f'<div class="sync-note">Jira data synced {last_sync:%Y-%m-%d %H:%M} · '
        f'{len(tickets):,} tickets since {tickets["created_date"].min():%d %b %Y}</div>',
        unsafe_allow_html=True,
    )

# ---------------------------------------------------------------------------
# Period selection
# ---------------------------------------------------------------------------
today = ta.jira_today(tickets)  # Jira timezone, not the PC clock
earliest = today - timedelta(days=config.TREND_HISTORY_DAYS)
DAY, WEEK, CUSTOM = "Day", "Week (Sun – Sat)", "Custom range"

c_mode, c_pick, c_weeks = st.columns([1.6, 1.6, 1.2])
with c_mode:
    mode = st.segmented_control("View", [DAY, WEEK, CUSTOM], default=DAY, key="trend_mode") or DAY
with c_pick:
    if mode == DAY:
        day = st.date_input("Date", value=today, min_value=earliest, max_value=today, format="DD/MM/YYYY")
        start = end = day
        period_label = f"{day:%d %b %Y}" + (" (today)" if day == today else "")
        prev_range = ta.previous_period(start, end)
        prev_label = "previous day"
    elif mode == WEEK:
        weeks_back = [ta.week_start(today) - timedelta(weeks=i) for i in range(config.TREND_HISTORY_DAYS // 7)]
        start = st.selectbox("Week", weeks_back, format_func=lambda d: ta.week_label(d)
                             + (" (this week)" if d == ta.week_start(today) else ""))
        end = min(start + timedelta(days=6), today)
        period_label = ta.week_label(start)
        prev_range = ta.previous_period(start, end, week=True)
        prev_label = ("previous week" if (end - start).days == 6
                      else f"same days last week ({prev_range[0]:%d %b} – {prev_range[1]:%d %b})")
    else:
        picked = st.date_input("Date range", value=(today - timedelta(days=6), today),
                               min_value=earliest, max_value=today, format="DD/MM/YYYY")
        start, end = (picked[0], picked[-1]) if isinstance(picked, (tuple, list)) and picked else (today, today)
        period_label = f"{start:%d %b %Y} – {end:%d %b %Y}"
        prev_range = ta.previous_period(start, end)
        prev_label = f"previous {(end - start).days + 1} days"
with c_weeks:
    trend_weeks = st.slider("Weekly trend: weeks shown", 4, config.TREND_HISTORY_DAYS // 7, 12)


# ---------------------------------------------------------------------------
# Analysis for the selected period
# ---------------------------------------------------------------------------
@st.cache_data(show_spinner=False)
def analyse(sync_stamp, start, end, prev_range):
    results = {}
    for cat in ta.CATEGORIES:
        top = ta.top_drivers(tickets, cat, start, end, prev_range)
        for d in top["drivers"]:
            summary = ta.driver_summary(tickets, cat, d["driver"], start, end)
            d["summary"] = summary["text"]
            d["tickets"] = summary["tickets"]
            d["first_seen"] = summary["all_tickets"]["created_date"].min()
            d["source"] = ta.driver_source_of(summary["all_tickets"])
            d["fixes"] = ta.match_fixes(tickets, fixes, cat, d["driver"], end)
        results[cat] = top
    return results


results = analyse(last_sync.isoformat(), start, end, prev_range)


def delta_html(cur, prev):
    if prev == 0:
        return '<span class="delta-flat">new</span>' if cur else '<span class="delta-flat">—</span>'
    pct = (cur - prev) / prev * 100
    cls = "delta-up" if pct > 0 else "delta-down" if pct < 0 else "delta-flat"
    arrow = "▲" if pct > 0 else "▼" if pct < 0 else "■"
    return f'<span class="{cls}">{arrow} {abs(pct):.0f}%</span>'


st.markdown(f'<div class="section-title">Overview &mdash; {html.escape(period_label)}</div>', unsafe_allow_html=True)
kpi_cols = st.columns(len(ta.CATEGORIES))
all_alerts = []
for col, cat in zip(kpi_cols, ta.CATEGORIES):
    r = results[cat]
    alerts = [(cat, d, f) for d in r["drivers"] for f in d["fixes"] if f["alert"]]
    all_alerts.extend(alerts)
    recurring = sum(1 for *_, f in alerts if f["alert"] == "recurring")
    col.markdown(
        f'<div class="kpi" style="--cat:{CATEGORY_COLOR[cat]};"><div class="kpi-label">{cat} tickets</div>'
        f'<div class="kpi-value">{r["total"]:,}</div>'
        f'<div class="kpi-sub">{delta_html(r["total"], r["prev_total"])} vs {prev_label} ({r["prev_total"]:,})'
        f' · top {len(r["drivers"])} drivers = {sum(d["count"] for d in r["drivers"]):,}'
        f' · {recurring} recurring-after-fix alert(s)</div></div>',
        unsafe_allow_html=True,
    )

if all_alerts:
    st.markdown('<div class="section-title">Fix alerts on the top drivers</div>', unsafe_allow_html=True)
    order = {"recurring": 0, "overdue": 1}
    for cat, d, f in sorted(all_alerts, key=lambda a: (order[a[2]["alert"]], -a[2]["since_count"])):
        tag = "Still occurring after fix" if f["alert"] == "recurring" else "Fix date passed"
        st.markdown(
            f'<div class="alert-box alert-{f["alert"]}"><span class="alert-tag">{tag}</span>'
            f'<span><b>{cat}</b> · {html.escape(d["driver"])} &rarr; '
            f'<a class="fix-key" href="{jira_url(f["key"])}" target="_blank">{f["key"]}</a>: '
            f'{html.escape(f["alert_text"])}</span></div>',
            unsafe_allow_html=True,
        )


# ---------------------------------------------------------------------------
# Charts
# ---------------------------------------------------------------------------
def short(text, n=60):
    return text if len(text) <= n else text[: n - 1] + "…"


def top_bar_chart(drivers, color):
    data = pd.DataFrame([{"Driver": short(d["driver"]), "Full driver": d["driver"], "Tickets": d["count"],
                          "Share": f'{d["share"]:.0f}%', "Previous period": d["prev"]} for d in drivers])
    base = alt.Chart(data).encode(
        y=alt.Y("Driver:N", sort="-x", title=None, axis=alt.Axis(labelLimit=380, labelColor=TEXT_SECONDARY)),
        x=alt.X("Tickets:Q", title=None, axis=alt.Axis(grid=True, gridColor="rgba(255,255,255,0.06)",
                                                         labelColor=TEXT_MUTED, tickMinStep=1)),
        tooltip=["Full driver", "Tickets", "Share", "Previous period"],
    )
    bars = base.mark_bar(color=color, cornerRadiusEnd=4, height=18)
    labels = base.mark_text(align="left", dx=6, color=TEXT_PRIMARY, fontWeight="bold").encode(text="Tickets:Q")
    return (bars + labels).properties(height=max(120, 44 * len(drivers))).configure_view(stroke=None)\
        .configure(background="transparent")


def weekly_chart(trend, top):
    """Top drivers only: "Other" usually dwarfs them, so it stays in the table."""
    colors = SERIES_COLORS[: len(top)]
    trend = trend[trend["driver"] != "Other"]
    trend = trend.assign(Driver=trend["driver"].map(lambda d: short(d, 45)),
                         Week=trend["week"].map(lambda w: f"{w:%d %b}"))
    weeks = [f"{w:%d %b}" for w in sorted(trend["week"].unique())]
    domain = [short(d, 45) for d in top]
    hover = alt.selection_point(fields=["Week"], nearest=True, on="pointerover", empty=False)
    base = alt.Chart(trend).encode(
        x=alt.X("Week:O", sort=weeks, title="Week starting (Sunday)",
                axis=alt.Axis(labelAngle=0, labelColor=TEXT_MUTED, titleColor=TEXT_MUTED, grid=False)),
        y=alt.Y("count:Q", title="Tickets", axis=alt.Axis(gridColor="rgba(255,255,255,0.06)",
                                                          labelColor=TEXT_MUTED, titleColor=TEXT_MUTED)),
        color=alt.Color("Driver:N", scale=alt.Scale(domain=domain, range=colors),
                        legend=alt.Legend(orient="bottom", columns=2, labelLimit=420, title=None,
                                          labelColor=TEXT_SECONDARY, symbolStrokeWidth=3)),
    )
    lines = base.mark_line(strokeWidth=2, interpolate="monotone")
    points = base.mark_point(filled=True, size=70, stroke=SURFACE_CARD, strokeWidth=2).encode(
        opacity=alt.condition(hover, alt.value(1), alt.value(0)),
        tooltip=[alt.Tooltip("week_label:N", title="Week"), alt.Tooltip("driver:N", title="Driver"),
                 alt.Tooltip("count:Q", title="Tickets")],
    ).add_params(hover)
    rule = alt.Chart(trend).mark_rule(color="rgba(255,255,255,0.25)").encode(
        x=alt.X("Week:O", sort=weeks)).transform_filter(hover)
    return (lines + points + rule).properties(height=320).configure_view(stroke=None)\
        .configure(background="transparent")


def fix_rows_html(fix_list):
    out = []
    for project in ("SAL", "SFS"):
        items = [f for f in fix_list if f["project"] == project]
        if not items:
            out.append(f'<div class="fix-row"><span class="fix-proj">{project}</span>'
                       f'<span class="fix-none">No {project} ticket in place</span></div>')
            continue
        for f in items:
            date_txt = ""
            if f["date"]:
                verb = {"deployed": "", "resolved": " · Resolved"}.get(f["state"], " · Planned")
                date_txt = f'{verb} {f["date"]:%d %b %Y}'
            elif f["state"] in ("in_progress", "not_started"):
                date_txt = " · No deployment date planned"
            match = (f'Linked to {f["linked_tickets"]} ticket(s)' if f["linked_tickets"]
                     else f'Linked via {f["via"]}' if f["via"] else f["match"])
            alert = (f'<div class="fix-alert {f["alert"]}">⚠ {html.escape(f["alert_text"])}</div>'
                     if f["alert"] else "")
            out.append(
                f'<div class="fix-row"><span class="fix-proj">{project}</span>'
                f'<a class="fix-key" href="{jira_url(f["key"])}" target="_blank" rel="noopener">{f["key"]}</a>'
                f'<span class="fix-chip" style="background:{FIX_STATE_COLOR[f["state"]]};">'
                f'{html.escape(f["label"])}{date_txt}</span>'
                f'<span class="fix-match">{match}</span>'
                f'<span class="fix-summary" title="{html.escape(f["summary"], quote=True)}">{html.escape(f["summary"])}</span>'
                f'{alert}</div>'
            )
    return "".join(out)


# ---------------------------------------------------------------------------
# Per-category tabs
# ---------------------------------------------------------------------------
st.markdown(f'<div class="section-title">Top {config.TREND_TOP_N} drivers &mdash; {html.escape(period_label)}</div>',
            unsafe_allow_html=True)
weekly = {}
tabs = st.tabs([f"{cat}  ({results[cat]['total']:,})" for cat in ta.CATEGORIES])
for tab, cat in zip(tabs, ta.CATEGORIES):
    r = results[cat]
    with tab:
        if not r["drivers"]:
            st.info(f"No {cat} tickets created in {period_label}.")
        else:
            if r["uncategorised"]:
                st.caption(f"{r['uncategorised']} ticket(s) with no category and no recognisable summary are "
                           "counted in the total but not ranked.")
            st.altair_chart(top_bar_chart(r["drivers"], CATEGORY_COLOR[cat]), width="stretch")
            for d in r["drivers"]:
                first_seen = f'{d["first_seen"]:%d %b %Y}' if pd.notna(d["first_seen"]) else "—"
                st.markdown(
                    f'<div class="driver-card"><div class="driver-head">'
                    f'<span class="driver-rank">#{d["rank"]}</span>'
                    f'<span class="driver-name">{html.escape(d["driver"])}</span>'
                    f'<span class="driver-count">{d["count"]:,}</span></div>'
                    f'<div class="driver-meta">{d["share"]:.0f}% of {cat} tickets · {delta_html(d["count"], d["prev"])} '
                    f'vs {prev_label} ({d["prev"]}) · grouped by {html.escape(d["source"])} · first seen {first_seen}</div>'
                    f'<div class="driver-summary">{html.escape(d["summary"])}</div>'
                    f'{fix_rows_html(d["fixes"])}</div>',
                    unsafe_allow_html=True,
                )
                with st.expander(f'Tickets for #{d["rank"]} ({d["count"]})'):
                    t = d["tickets"]
                    cols = {"key": "Key", "created": "Created", "status": "Status", "sub_team": "Team / Account",
                            "severity": "Severity", "summary": "Summary"}
                    if cat == "Order Fallout":
                        cols.update({"of_reason": "Fallout Reason", "of_step": "Failed Step", "of_error": "Error"})
                    else:
                        cols.update({"inc_cat": "Incident Category", "inc_subcat": "Incident Sub Category"})
                    view = t[list(cols)].rename(columns=cols)
                    view["Key"] = view["Key"].map(jira_url)
                    view["Created"] = view["Created"].str[:16].str.replace("T", " ")
                    view["Linked"] = t["links"].map(", ".join).values
                    st.dataframe(view, hide_index=True, width="stretch", height=min(400, 40 + 35 * len(view)),
                                 column_config={"Key": st.column_config.LinkColumn("Key", display_text=r".*/browse/(.*)$")})

        st.markdown(f'<div class="section-title">{cat} &mdash; weekly trend (Sunday – Saturday), '
                    f'last {trend_weeks} weeks</div>', unsafe_allow_html=True)
        trend, top = ta.weekly_trend(tickets, cat, trend_weeks, end)
        weekly[cat] = trend
        if trend.empty:
            st.info("No tickets in this window.")
        else:
            st.altair_chart(weekly_chart(trend, top), width="stretch")
            st.caption(f"Top {len(top)} drivers of the window; all other drivers are in the table "
                       f"(\"Other\" row). The week of {ta.week_label(ta.week_start(end))} runs to "
                       f"{end:%d %b}" + (" (in progress)." if end == today else "."))
            with st.expander("Weekly trend as a table"):
                pivot = trend.pivot_table(index="driver", columns="week_label", values="count", sort=False)
                pivot = pivot[[ta.week_label(w) for w in sorted(trend["week"].unique())]]
                st.dataframe(pivot, width="stretch")


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------
def build_excel():
    drivers_rows, ticket_frames, trend_frames = [], [], []
    for cat in ta.CATEGORIES:
        for d in results[cat]["drivers"]:
            alerts = "; ".join(f"{f['key']}: {f['alert_text']}" for f in d["fixes"] if f["alert"])
            drivers_rows.append({
                "Period": period_label, "Category": cat, "Rank": d["rank"], "Driver": d["driver"],
                "Tickets": d["count"], "Share %": d["share"], "Compared with": prev_label,
                "Tickets then": d["prev"],
                "Grouped by": d["source"], "Summary": d["summary"],
                "SAL in place": ta.fix_status_text(d["fixes"], "SAL"),
                "SFS in place": ta.fix_status_text(d["fixes"], "SFS"),
                "Alert": alerts,
            })
            t = d["tickets"].assign(Category=cat, Driver=d["driver"], Rank=d["rank"])
            ticket_frames.append(t)
        if cat in weekly and not weekly[cat].empty:
            trend_frames.append(weekly[cat].assign(Category=cat))

    fix_rows = []
    for cat in ta.CATEGORIES:
        for d in results[cat]["drivers"]:
            for f in d["fixes"]:
                fix_rows.append({"Category": cat, "Rank": d["rank"], "Driver": d["driver"], "Ticket": f["key"],
                                 "Project": f["project"], "Match": f["match"], "Jira status": f["status"],
                                 "Fix state": f["label"], "Date": f["date"].isoformat() if f["date"] else "",
                                 "Alert": f["alert_text"], "Summary": f["summary"], "Link": jira_url(f["key"])})

    sheets = {"Top Drivers": pd.DataFrame(drivers_rows), "Fix Tracking (SAL-SFS)": pd.DataFrame(fix_rows)}
    if trend_frames:
        tr = pd.concat(trend_frames)
        sheets["Weekly Trend"] = tr.pivot_table(index=["Category", "driver"], columns="week", values="count",
                                                sort=False).rename(columns=ta.week_label).reset_index()\
            .rename(columns={"driver": "Driver"})
    if ticket_frames:
        tk = pd.concat(ticket_frames)
        sheets["Tickets"] = pd.DataFrame({
            "Category": tk["Category"], "Rank": tk["Rank"], "Driver": tk["Driver"], "Key": tk["key"],
            "Created": tk["created"].str[:16].str.replace("T", " "), "Status": tk["status"],
            "Team / Account": tk["sub_team"], "Severity": tk["severity"], "Summary": tk["summary"],
            "Incident Sub Category": tk["inc_subcat"], "Fallout Reason": tk["of_reason"],
            "Failed Step": tk["of_step"], "Error": tk["of_error"],
            "Linked": tk["links"].map(", ".join), "Link": tk["key"].map(jira_url),
        })

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            header_fill = PatternFill("solid", fgColor="20202C")
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
            for idx, col_name in enumerate(frame.columns, start=1):
                width = min(70, max(len(str(col_name)), frame[col_name].astype(str).str.len().max() or 0) + 2)
                ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.freeze_panes = "A2"
    return buf.getvalue()


st.markdown('<div class="section-title">Download</div>', unsafe_allow_html=True)
st.download_button(
    "⬇️ Download trend analysis (Excel)",
    data=build_excel(),
    file_name=f"Ticket_Trend_Analysis_{start:%Y-%m-%d}_to_{end:%Y-%m-%d}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    help="Top drivers with summaries, SAL / SFS fix tracking and alerts, weekly trend and the ticket list.",
)
