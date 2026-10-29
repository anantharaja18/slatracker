"""
Streamlit dashboard: live view of all 4 ticket categories with breach status.

Run with:
    streamlit run dashboard.py
"""

import html
import itertools
import io
from collections import defaultdict
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl.styles import Font
from streamlit_autorefresh import st_autorefresh

# import ai_assistant  # AI chat assistant -- disabled until ANTHROPIC_API_KEY is set up
import config
import jira_client
import breach_logic
import ui_table
# Jira status dot colors ("Tickets by Status" cards, popups) and severity tag
# colors are shared with the styled tables.
from ui_table import DEFAULT_STATUS_DOT_COLOR, SEVERITY_TAG_COLORS, STATUS_DOT_COLOR


# ---------------------------------------------------------------------------
# Palette (validated categorical + fixed status colors, dark-surface safe)
# ---------------------------------------------------------------------------
SURFACE_PAGE = "#0d0d14"
SURFACE_CARD = "#14141c"
SURFACE_PANEL = "#1a1a24"
TEXT_PRIMARY = "#F5F5F7"
TEXT_SECONDARY = "#C3C2B7"
TEXT_MUTED = "#898781"
ACCENT_FROM = "#7C3AED"
ACCENT_TO = "#DB2777"
HEADER_BG = "#20202c"

COLOR_GOOD = "#0ca30c"
COLOR_CAUTION = "#7cb342"
COLOR_WARNING = "#fab219"
COLOR_SERIOUS = "#ec835a"
COLOR_CRITICAL = "#d03b3b"
COLOR_CRITICAL_DARK = "#7a1f1f"
COLOR_INFLOW = "#3987e5"   # tickets created today
COLOR_BACKLOG = "#8b7fd6"  # tickets created before today

STATUS_ORDER = [
    breach_logic.STATUS_OK,
    breach_logic.STATUS_WARNING,
    breach_logic.STATUS_BREACHED,
    breach_logic.STATUS_UNKNOWN,
]
STATUS_COLOR = {
    breach_logic.STATUS_OK: COLOR_GOOD,
    breach_logic.STATUS_WARNING: COLOR_WARNING,
    breach_logic.STATUS_BREACHED: COLOR_CRITICAL,
    breach_logic.STATUS_UNKNOWN: "#898781",
}

# Time-to-resolve buckets: every open ticket falls into exactly one, so a
# category's bucket counts always add up to its total open-ticket count.
BUCKET_BREACHED_PAST = "Breached (Past)"
BUCKET_BREACHED_TODAY = "Breached Today"
BUCKET_15 = "Within 15 mins"
BUCKET_30 = "Within 30 mins"
BUCKET_45 = "Within 45 mins"
BUCKET_OVER_1H = "> 1 hour"
BUCKET_ORDER = [BUCKET_BREACHED_PAST, BUCKET_BREACHED_TODAY, BUCKET_15, BUCKET_30, BUCKET_45, BUCKET_OVER_1H]
BUCKET_COLOR = {
    BUCKET_BREACHED_PAST: COLOR_CRITICAL_DARK,
    BUCKET_BREACHED_TODAY: COLOR_CRITICAL,
    BUCKET_15: COLOR_SERIOUS,
    BUCKET_30: COLOR_WARNING,
    BUCKET_45: COLOR_CAUTION,
    BUCKET_OVER_1H: COLOR_GOOD,
}

CATEGORY_COLOR = {
    "Retail": "#3987e5",
    "SAC": "#d95926",
    "B2B": "#199e70",
    "Order Fallout": "#c98500",
}
DEFAULT_CATEGORY_COLOR = "#3987e5"

def status_dot_color(status):
    return STATUS_DOT_COLOR.get((status or "").strip().lower(), DEFAULT_STATUS_DOT_COLOR)


def compute_status_counts(df):
    if df.empty:
        return {}
    counts = df["Status"].value_counts()
    return dict(sorted(counts.items(), key=lambda kv: kv[1], reverse=True))


def open_tickets_with_status(df, status):
    if df.empty:
        return []
    matched = df[df["Status"].str.strip().str.lower() == status]
    return tickets_list(matched)


def bucket_for(hours_remaining, due_dt):
    if hours_remaining is None or pd.isna(hours_remaining) or due_dt is None:
        return BUCKET_OVER_1H
    minutes = hours_remaining * 60
    if minutes < 0:
        jira_tz = due_dt.tzinfo
        if due_dt.date() < datetime.now(jira_tz).date():
            return BUCKET_BREACHED_PAST
        return BUCKET_BREACHED_TODAY
    if minutes <= 15:
        return BUCKET_15
    if minutes <= 30:
        return BUCKET_30
    if minutes <= 45:
        return BUCKET_45
    return BUCKET_OVER_1H


def compute_time_buckets(df):
    buckets = {b: {"count": 0, "tickets": []} for b in BUCKET_ORDER}
    if df.empty:
        return buckets
    for _, row in df.iterrows():
        b = bucket_for(row["Hours Remaining"], row["Due Datetime"])
        buckets[b]["count"] += 1
        buckets[b]["tickets"].append(popup_ticket(row))
    return buckets


def jira_ticket_url(key):
    return f'{config.JIRA_BASE_URL.rstrip("/")}/browse/{key}'


def bucket_popup_html(tickets):
    if not tickets:
        return '<div class="bucket-popup-inner"><div class="bucket-popup-empty">No tickets</div></div>'
    # data-* attributes feed the popup's CSV download (see POPUP_SCRIPT).
    items = "".join(
        f'<a class="bucket-popup-item" href="{html.escape(jira_ticket_url(t["key"]), quote=True)}" '
        f'target="_blank" rel="noopener noreferrer" '
        f'data-key="{html.escape(t["key"], quote=True)}" '
        f'data-severity="{html.escape(t.get("severity") or "", quote=True)}" '
        f'data-assignee="{html.escape(t["assignee"], quote=True)}" '
        f'data-status="{html.escape(t.get("status") or "", quote=True)}" '
        f'data-due="{html.escape(t.get("due") or "", quote=True)}" '
        f'data-summary="{html.escape(t.get("summary") or "", quote=True)}">'
        f'<span class="bucket-popup-key">{html.escape(t["key"])}</span>'
        f'{severity_tag_html(t.get("severity"))}'
        f'{status_tag_html(t.get("status"))}'
        f'<span class="bucket-popup-assignee" title="{html.escape(t["assignee"])}">{html.escape(t["assignee"])}</span></a>'
        for t in tickets
    )
    return f'<div class="bucket-popup-inner">{items}</div>'


def status_tag_html(status):
    """Current Jira status (Escalated, In Progress, ...) as a pill, with the
    same dot color as the Tickets by Status cards."""
    if not status:
        return ""
    return (
        f'<span class="bucket-popup-status"><span class="dot" style="background:{status_dot_color(status)};"></span>'
        f'{html.escape(status)}</span>'
    )


def severity_tag_html(severity):
    if not severity:
        return ""
    bg, fg = SEVERITY_TAG_COLORS.get(severity, (DEFAULT_STATUS_DOT_COLOR, TEXT_PRIMARY))
    return f'<span class="bucket-popup-tag" style="background:{bg};color:{fg};">{html.escape(severity)}</span>'


_popup_counter = itertools.count()

# Export groups for the section-level "Download breached" buttons.
TTR_BREACHED_GROUP = "ttr-breached"            # Time to Resolve: Breached (Past) + Breached Today
RESOLVED_BREACHED_GROUP = "resolved-breached"  # Today's inflow: Resolved late + Open, breached

DOWNLOAD_ICON_SVG = (
    '<svg viewBox="0 0 24 24" width="15" height="15" fill="none" stroke="currentColor" '
    'stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round">'
    '<path d="M12 4v11"/><path d="M7 10l5 5 5-5"/><path d="M5 20h14"/></svg>'
)


def popup_row(row_inner_html, popup_inner_html, title, row_class="bucket-row", export_group=None):
    """A bucket row (or other element, via row_class) whose ticket popup opens
    on hover and closes a few seconds after the pointer leaves it (or via its
    x). Handled by POPUP_SCRIPT. Popups sharing an export_group are downloaded
    together by a section_title_with_download() button."""
    popup_id = f"popup-{next(_popup_counter)}"
    group_attr = f' data-export-group="{html.escape(export_group, quote=True)}"' if export_group else ""
    return (
        f'<div class="popup-wrap" data-popup-id="{popup_id}"{group_attr}>'
        f'<div class="{row_class} has-popup">{row_inner_html}</div>'
        f'<div class="bucket-popup"><div class="bucket-popup-box">'
        f'<div class="bucket-popup-header"><span class="bucket-popup-title">{html.escape(title)}</span>'
        f'<span class="bucket-popup-actions">'
        f'<span class="bucket-popup-download" title="Download this list (CSV)">{DOWNLOAD_ICON_SVG}</span>'
        f'<span class="bucket-popup-close" title="Close">&times;</span></span></div>'
        f'{popup_inner_html}</div></div>'
        f'</div>'
    )


def section_title_with_download(title_html, export_group, button_label, file_name):
    """A section title bar with a button on the right that downloads every
    popup in export_group as one CSV."""
    return (
        f'<div class="section-title section-title-bar"><span>{title_html}</span>'
        f'<span class="section-download" data-export-group="{html.escape(export_group, quote=True)}" '
        f'data-file-name="{html.escape(file_name, quote=True)}" title="Download as CSV">'
        f'{DOWNLOAD_ICON_SVG}<span>{html.escape(button_label)}</span></span></div>'
    )


# Hovering a row briefly opens its popup (closing any other open one); it
# closes 5 seconds after the pointer leaves the row and popup, or at once when
# its x is clicked. The open popup's id is remembered so it's
# re-opened after a Streamlit rerun (e.g. the 1-minute auto-refresh) redraws it.
POPUP_SCRIPT = """
<script>
(function () {
  const win = window.parent && window.parent.document ? window.parent : window;
  const doc = win.document;
  if (win.__slaPopupsInstalled) return;
  win.__slaPopupsInstalled = true;
  win.__slaOpenPopup = null;

  function setOpen(id) {
    win.__slaOpenPopup = id;
    doc.querySelectorAll(".popup-wrap.open").forEach(function (w) {
      if (w.dataset.popupId !== id) w.classList.remove("open");
    });
    if (id) {
      const w = doc.querySelector('.popup-wrap[data-popup-id="' + id + '"]');
      if (w) openWrap(w);
    }
  }

  // Shows a popup, flipping it to align with the row's right edge if it
  // would otherwise run past the right side of the window.
  function openWrap(w) {
    w.classList.add("open");
    const pop = w.querySelector(".bucket-popup");
    if (!pop) return;
    pop.classList.remove("align-right");
    if (pop.getBoundingClientRect().right > doc.documentElement.clientWidth - 8) {
      pop.classList.add("align-right");
    }
  }

  // Row left under the pointer when a popup is closed; it must not instantly
  // re-open, so it's ignored until the pointer moves off it.
  let suppressRow = null;

  // A row must be hovered briefly before its popup opens, so sweeping the
  // mouse across the page doesn't flash popups. Once the pointer leaves the
  // open row/popup, it closes after a grace period unless the pointer returns.
  const OPEN_DELAY_MS = 500;
  const CLOSE_DELAY_MS = 5000;
  let openTimer = null, pendingId = null, closeTimer = null;

  function cancelOpen() {
    clearTimeout(openTimer);
    openTimer = null;
    pendingId = null;
  }
  function cancelClose() {
    clearTimeout(closeTimer);
    closeTimer = null;
  }
  function scheduleClose() {
    if (!win.__slaOpenPopup || closeTimer) return;
    closeTimer = setTimeout(function () {
      closeTimer = null;
      setOpen(null);
    }, CLOSE_DELAY_MS);
  }

  doc.addEventListener("mouseover", function (e) {
    const wrap = e.target.closest && e.target.closest(".popup-wrap");
    const row = e.target.closest && e.target.closest(".has-popup");
    if (suppressRow && row !== suppressRow) suppressRow = null;

    if (wrap && wrap.dataset.popupId === win.__slaOpenPopup) {
      cancelClose();
      cancelOpen();
      return;
    }
    scheduleClose();
    if (!row || row === suppressRow) {
      cancelOpen();
      return;
    }
    const id = row.closest(".popup-wrap").dataset.popupId;
    if (id === pendingId) return;
    cancelOpen();
    pendingId = id;
    openTimer = setTimeout(function () {
      openTimer = null;
      pendingId = null;
      cancelClose();
      setOpen(id);
    }, OPEN_DELAY_MS);
  });

  // Pointer left the window entirely.
  doc.addEventListener("mouseout", function (e) {
    if (!e.relatedTarget) {
      cancelOpen();
      scheduleClose();
    }
  });

  // Saves a popup's ticket list as CSV, built from the popup's own rows so it
  // always matches the count that was hovered.
  function csvCell(v) {
    return '"' + String(v == null ? "" : v).replace(/"/g, '""') + '"';
  }
  // One or more popups into one CSV. Popup titles are "Team, em dash, List".
  // (No angle brackets in this script: Streamlit's HTML handling drops it.)
  const TITLE_SEP = " " + String.fromCharCode(8212) + " ";
  function downloadPopups(wraps, fileBase) {
    const lines = [["S.No", "Team", "List", "Key", "Severity", "Status", "Assignee", "Due", "Summary", "Jira Link"]];
    wraps.forEach(function (wrap) {
      const titleEl = wrap.querySelector(".bucket-popup-title");
      const parts = (titleEl ? titleEl.textContent.trim() : "").split(TITLE_SEP);
      const team = parts.length > 1 ? parts[0] : "";
      const list = parts.length > 1 ? parts.slice(1).join(TITLE_SEP) : parts[0];
      wrap.querySelectorAll(".bucket-popup-item").forEach(function (a) {
        const d = a.dataset;
        const key = d.key || (a.querySelector(".bucket-popup-key") || a).textContent.trim();
        lines.push([lines.length, team, list, key, d.severity, d.status, d.assignee, d.due, d.summary, a.href]);
      });
    });
    if (lines.length < 2) return;
    const csv = "\\ufeff" + lines.map(function (r) { return r.map(csvCell).join(","); }).join("\\r\\n");
    const blob = new Blob([csv], { type: "text/csv;charset=utf-8" });
    const url = URL.createObjectURL(blob);
    const link = doc.createElement("a");
    const now = new Date();  // local date, not UTC, so it matches the dashboard's "today"
    const stamp = now.getFullYear() + "-" + String(now.getMonth() + 1).padStart(2, "0") + "-" + String(now.getDate()).padStart(2, "0");
    link.href = url;
    link.download = fileBase.replace(/[^A-Za-z0-9]+/g, "_").replace(/^_|_$/g, "") + "_" + stamp + ".csv";
    doc.body.appendChild(link);
    link.click();
    link.remove();
    setTimeout(function () { URL.revokeObjectURL(url); }, 1000);
  }

  doc.addEventListener("click", function (e) {
    const dl = e.target.closest && e.target.closest(".bucket-popup-download");
    if (dl) {
      e.preventDefault();
      const wrap = dl.closest(".popup-wrap");
      const titleEl = wrap.querySelector(".bucket-popup-title");
      downloadPopups([wrap], titleEl ? titleEl.textContent.trim() : "tickets");
      return;
    }
    // Section-level download: every popup tagged with the same export group.
    const sdl = e.target.closest && e.target.closest(".section-download");
    if (sdl) {
      e.preventDefault();
      const group = sdl.dataset.exportGroup;
      downloadPopups(doc.querySelectorAll('.popup-wrap[data-export-group="' + group + '"]'), sdl.dataset.fileName || group);
      return;
    }
    if (e.target.closest && e.target.closest(".bucket-popup-close")) {
      e.preventDefault();
      cancelOpen();
      cancelClose();
      setOpen(null);
      const under = doc.elementFromPoint(e.clientX, e.clientY);
      suppressRow = under && under.closest ? under.closest(".has-popup") : null;
    }
  });

  new win.MutationObserver(function () {
    const id = win.__slaOpenPopup;
    if (!id) return;
    const w = doc.querySelector('.popup-wrap[data-popup-id="' + id + '"]');
    if (w && !w.classList.contains("open")) openWrap(w);
  }).observe(doc.body, { childList: true, subtree: true });
})();
</script>
"""


def escalated_tickets_with_links(df):
    if df.empty:
        return []
    escalated = df[df["Status"].str.strip().str.lower() == "escalated"]
    return [{"key": row["Key"], "links": row["Linked Issues"]} for _, row in escalated.iterrows()]


def ticket_link_html(key, css_class="esc-key", title=""):
    title_attr = f' title="{html.escape(title, quote=True)}"' if title else ""
    return (
        f'<a class="{css_class}" href="{html.escape(jira_ticket_url(key), quote=True)}" '
        f'target="_blank" rel="noopener noreferrer"{title_attr}>{html.escape(key)}</a>'
    )


def ticket_details_view(df):
    """Ticket Details table: S.No from 1, readable Created time, internal
    columns dropped (ui_table links the Key and styles status / SLA cells)."""
    def fmt_created(raw):
        dt = breach_logic._parse_datetime(raw)
        return dt.strftime("%Y-%m-%d %H:%M") if dt else ""

    view = pd.DataFrame({
        "S.No": range(1, len(df) + 1),
        "Key": df["Key"],
        "Summary": df["Summary"],
        "Status": df["Status"],
        "Priority": df["Priority"],
        "Severity": df["Severity"],
        "Assignee": df["Assignee"],
        "Created": df["Created"].map(fmt_created),
        "Due": df["Due"],
        "Hours Remaining": df["Hours Remaining"],
        "Breach Status": df["Breach Status"],
        "Escalation Path": df["Escalation Path"],
        "Linked Tickets": df["Linked Issues"].map(lambda links: ", ".join(l["key"] for l in links or [])),
    })
    return view.reset_index(drop=True)


def escalated_table_html(df):
    """Escalated tickets as an HTML table (st.dataframe can't hold several links
    in one cell): S.No from 1, the ticket key and each linked ticket open in
    Jira, and hovering a linked ticket shows its summary."""
    columns = ["S.No", "Key", "Category", "Summary", "Severity", "Escalation Path", "Linked Tickets", "Assignee", "Due"]
    head = "".join(f"<th>{c}</th>" for c in columns)
    body = []
    for n, (_, row) in enumerate(df.iterrows(), start=1):
        links = row["Linked Issues"] or []
        linked_html = (
            "".join(ticket_link_html(l["key"], "esc-link", l["summary"]) for l in links)
            if links else '<span class="esc-none">—</span>'
        )
        cells = [
            f'<td class="esc-num">{n}</td>',
            f'<td>{ticket_link_html(row["Key"])}</td>',
            f'<td>{html.escape(row["Category"])}</td>',
            f'<td class="esc-summary" title="{html.escape(row["Summary"], quote=True)}">{html.escape(row["Summary"])}</td>',
            f'<td>{html.escape(row["Severity"])}</td>',
            f'<td>{html.escape(row["Escalation Path"] or "—")}</td>',
            f'<td class="esc-links">{linked_html}</td>',
            f'<td>{html.escape(row["Assignee"])}</td>',
            f'<td class="esc-due">{html.escape(row["Due"])}</td>',
        ]
        body.append(f'<tr>{"".join(cells)}</tr>')
    return (
        f'<div class="esc-table-wrap"><table class="esc-table"><thead><tr>{head}</tr></thead>'
        f'<tbody>{"".join(body)}</tbody></table></div>'
    )


def escalated_links_popup_html(tickets):
    if not tickets:
        return '<div class="bucket-popup-inner"><div class="bucket-popup-empty">No escalated tickets</div></div>'
    blocks = []
    for t in tickets:
        link_count = len(t["links"])
        summary_text = f'{link_count} linked ticket{"s" if link_count != 1 else ""}' if link_count else "No linked tickets"
        blocks.append(
            f'<a class="bucket-popup-item" href="{html.escape(jira_ticket_url(t["key"]), quote=True)}" '
            f'target="_blank" rel="noopener noreferrer">'
            f'<span class="bucket-popup-key">{html.escape(t["key"])}</span>'
            f'<span class="bucket-popup-assignee" title="{html.escape(summary_text)}">{html.escape(summary_text)}</span></a>'
        )
        blocks.extend(
            f'<a class="bucket-popup-item bucket-popup-sublink" href="{html.escape(jira_ticket_url(l["key"]), quote=True)}" '
            f'target="_blank" rel="noopener noreferrer">'
            f'<span class="bucket-popup-key">&#8618; {html.escape(l["key"])}</span>'
            f'<span class="bucket-popup-assignee" title="{html.escape(l["summary"])}">{html.escape(l["summary"])}</span></a>'
            for l in t["links"]
        )
    return f'<div class="bucket-popup-inner">{"".join(blocks)}</div>'


APP_NAME = "JIRA - SLA Tracker & Notifier"
APP_SUBTITLE = "Real-time ticket monitoring across Retail, SAC, B2B and Order Fallout queues"


def bell_svg(size):
    return (
        f'<svg viewBox="0 0 24 24" width="{size}" height="{size}" fill="white">'
        '<path d="M12 22c1.1 0 2-.9 2-2h-4c0 1.1.89 2 2 2zm6-6v-5c0-3.07-1.64-5.64-4.5-6.32V4c0-.83-.67-1.5-1.5-1.5s-1.5.67-1.5 1.5v.68C7.63 5.36 6 7.92 6 11v5l-2 2v1h16v-1l-2-2zm-9.44-8.6l-1.06-1.06C6.09 5.64 5 8.51 5 11h1.5c0-2.42.72-4.78 2.06-6.6zM19.5 11H21c0-2.49-1.09-5.36-2.5-7.66l-1.06 1.06C18.78 6.22 19.5 8.58 19.5 11z"/>'
        '</svg>'
    )


HEADER_HTML = (
    f'<div class="app-header"><div class="app-logo">{bell_svg(19)}</div><div>'
    f'<div class="app-title">{html.escape(APP_NAME)}</div>'
    f'<div class="app-subtitle">{html.escape(APP_SUBTITLE)}</div>'
    '</div></div>'
)

# Shown on a session's first load while every team's tickets are fetched from
# Jira (~20 s): the same header, in the same place as the dashboard's, with
# the loading status below it. Later auto-refreshes keep the previous
# dashboard on screen instead.
LOADER_HTML = (
    f'<div class="app-loader">{HEADER_HTML}'
    '<div class="app-loader-status"><span class="app-loader-spinner"></span>'
    'Loading dashboard&hellip; fetching the latest tickets from Jira</div></div>'
)

st.set_page_config(page_title="Jira Breach Monitor", page_icon="🔔", layout="wide")

AUTO_REFRESH_SECONDS = 60  # matches load_category_data/load_resolved_data cache ttl
st_autorefresh(interval=AUTO_REFRESH_SECONDS * 1000, key="dashboard_autorefresh")

st.markdown(
    f"""
    <style>
        #MainMenu, footer, header {{visibility: hidden;}}
        /* Hidden header and invisible helper elements (auto-refresh timer, popup
           script, this style block) still take up space; collapse them so the
           title sits at the top of the card. */
        [data-testid="stHeader"] {{ display: none; }}
        [data-testid="stElementContainer"]:has(> iframe[height="0"]),
        [data-testid="stElementContainer"]:has(iframe[title="streamlit_autorefresh.st_autorefresh"]),
        [data-testid="stElementContainer"]:has([data-testid="stCustomComponentV1"]),
        [data-testid="stElementContainer"]:has(style),
        [data-testid="stElementContainer"]:has(script) {{ display: none; }}

        [data-testid="stAppViewContainer"] {{
            background:
                radial-gradient(circle at 0% 0%, rgba(124,58,237,0.22), transparent 45%),
                radial-gradient(circle at 100% 0%, rgba(31,111,235,0.18), transparent 45%),
                {SURFACE_PAGE};
        }}
        [data-testid="stMain"] {{ background: transparent; overflow-x: hidden; }}
        /* Keep the dashboard steady while Streamlit reruns (auto-refresh): no stale-element fade. */
        [data-stale="true"] {{ opacity: 1 !important; transition: none !important; }}
        html, body, [data-testid="stApp"], [data-testid="stAppViewContainer"] {{ overflow-x: hidden; }}

        .block-container, [data-testid="stMainBlockContainer"] {{
            background: {SURFACE_CARD};
            border: 1px solid rgba(255,255,255,0.07);
            border-radius: 20px;
            padding: 0.9rem 1.8rem 1.6rem 1.8rem !important;
            margin-top: 0.6rem;
            margin-bottom: 0.8rem;
            max-width: min(1880px, calc(100vw - 2.5rem));  /* use the screen width, keep a small side margin */
            box-shadow: 0 20px 60px rgba(0,0,0,0.45);
        }}
        [data-testid="stVerticalBlock"] {{ gap: 0.55rem; }}

        .app-header {{ display: flex; align-items: center; gap: 10px; }}
        .app-logo {{
            width: 36px; height: 36px; border-radius: 10px; flex-shrink: 0;
            background: linear-gradient(135deg, {ACCENT_FROM}, {ACCENT_TO});
            display: flex; align-items: center; justify-content: center; font-size: 17px;
            position: relative; animation: siren-pulse 2.2s ease-in-out infinite;
        }}
        @keyframes siren-pulse {{
            0%, 100% {{ box-shadow: 0 0 0 0 rgba(219,39,119,0.55); }}
            50% {{ box-shadow: 0 0 0 7px rgba(219,39,119,0); }}
        }}
        .app-logo svg {{ display: block; }}
        .app-title {{ font-size: 1.15rem; font-weight: 700; color: {TEXT_PRIMARY}; line-height: 1.1; }}
        .app-subtitle {{ font-size: 0.75rem; color: {TEXT_MUTED}; margin-top: 1px; }}
        .app-refreshed {{ font-size: 0.68rem; color: {TEXT_MUTED}; text-align: right; margin-top: 6px; line-height: 1.35; }}
        a.trend-link {{
            display: block; text-align: center; padding: 6px 10px; border-radius: 8px;
            background: linear-gradient(135deg, {ACCENT_FROM}, {ACCENT_TO});
            color: white !important; font-size: 0.85rem; font-weight: 600; text-decoration: none !important;
        }}
        a.trend-link:hover {{ filter: brightness(1.12); }}
        .page-links {{ display: flex; flex-direction: column; gap: 6px; }}
        .page-links a.trend-link {{ font-size: 0.8rem; padding: 5px 8px; }}
        .header-spacer {{ height: 2px; }}

        .app-loader {{ min-height: 60vh; }}
        .app-loader-status {{
            display: flex; align-items: center; gap: 10px; margin-top: 30px;
            font-size: 0.9rem; color: {TEXT_SECONDARY};
        }}
        .app-loader-spinner {{
            width: 18px; height: 18px; border-radius: 50%; display: inline-block;
            border: 2.5px solid rgba(255,255,255,0.15); border-top-color: {ACCENT_TO};
            animation: app-loader-spin 0.9s linear infinite;
        }}
        @keyframes app-loader-spin {{ to {{ transform: rotate(360deg); }} }}

        .overall-summary {{
            display: flex; align-items: center; justify-content: space-evenly; gap: 12px;
            height: 100%; padding: 10px 16px;
            background: rgba(255,255,255,0.04); border: 1px solid rgba(255,255,255,0.08);
            border-radius: 12px;
        }}
        .overall-stat {{ text-align: center; }}
        .overall-stat-value {{ font-size: 1.1rem; font-weight: 700; color: {TEXT_PRIMARY}; line-height: 1.1; }}
        .overall-stat-label {{
            font-size: 0.6rem; font-weight: 700; letter-spacing: 0.05em; text-transform: uppercase;
            color: {TEXT_MUTED}; margin-top: 2px; white-space: nowrap;
        }}
        .overall-divider {{ width: 1px; height: 28px; background: rgba(255,255,255,0.12); }}

        .section-title {{
            font-size: 0.72rem; font-weight: 700; letter-spacing: 0.06em; text-transform: uppercase;
            color: {TEXT_SECONDARY}; background: rgba(255,255,255,0.04);
            border: 1px solid rgba(255,255,255,0.08); border-radius: 8px;
            padding: 7px 14px; margin: 6px 0 2px 0;
        }}
        .section-title-bar {{ display: flex; align-items: center; justify-content: space-between; gap: 12px; }}
        .section-download {{
            display: inline-flex; align-items: center; gap: 6px; cursor: pointer; flex-shrink: 0;
            padding: 3px 10px; margin: -3px -6px -3px 0; border-radius: 6px;
            border: 1px solid rgba(255,255,255,0.14); color: {TEXT_SECONDARY};
            font-size: 0.66rem; letter-spacing: 0.04em;
        }}
        .section-download:hover {{ background: rgba(255,255,255,0.08); color: {TEXT_PRIMARY}; }}

        .bucket-card {{
            border: 1px solid rgba(255,255,255,0.08); border-radius: 12px;
            background: {SURFACE_PANEL};
        }}
        .bucket-card-head {{
            background: {HEADER_BG}; color: {TEXT_PRIMARY}; font-weight: 700; text-align: center;
            padding: 7px 8px 6px 8px; font-size: 0.82rem; border-radius: 12px 12px 0 0; min-height: 2.6rem;
            display: flex; align-items: center; justify-content: center;
            border-bottom: 3px solid transparent;
            border-image: linear-gradient(90deg, {ACCENT_FROM}, {ACCENT_TO}) 1;
        }}

        .summary-card {{ display: flex; flex-direction: column; min-height: 290px; }}
        .mini-stats {{
            display: flex; align-items: center; justify-content: space-evenly;
            padding: 10px 2px; background: rgba(255,255,255,0.03);
        }}
        .summary-card .mini-stats {{ flex: 1; }}  /* strips absorb any height difference between cards */
        .mini-stat {{ text-align: center; padding: 4px 3px; border-radius: 8px; min-width: 0; }}
        .mini-stat.has-popup {{ cursor: pointer; }}
        .mini-stat.has-popup:hover {{ background: rgba(255,255,255,0.06); }}
        .mini-stat-value {{ font-size: 1.3rem; font-weight: 700; color: {TEXT_PRIMARY}; line-height: 1.1; }}
        .mini-stats-grow .mini-stat-value {{ font-size: 1.9rem; }}
        .mini-stat-label {{
            font-size: 0.56rem; font-weight: 700; letter-spacing: 0.03em; text-transform: uppercase;
            color: {TEXT_MUTED}; margin-top: 3px; white-space: nowrap;
        }}
        .mini-stat-divider {{ width: 1px; align-self: stretch; margin: 6px 0; background: rgba(255,255,255,0.1); }}
        .bucket-subcaption {{
            font-size: 0.6rem; font-weight: 700; letter-spacing: 0.05em; text-transform: uppercase;
            color: {TEXT_MUTED}; padding: 8px 14px 4px 14px; border-top: 1px solid rgba(255,255,255,0.12);
        }}
        .bucket-row {{
            position: relative;
            display: flex; align-items: center; justify-content: space-between;
            padding: 7px 14px; border-top: 1px solid rgba(255,255,255,0.06);
        }}
        .bucket-row:last-child {{ border-radius: 0 0 12px 12px; }}
        .bucket-row .dot {{ width: 8px; height: 8px; border-radius: 50%; display: inline-block; margin-right: 8px; }}
        .bucket-row .label {{ color: {TEXT_SECONDARY}; font-size: 0.78rem; display: flex; align-items: center; cursor: default; }}
        .bucket-row .value {{ color: {TEXT_PRIMARY}; font-weight: 700; font-size: 0.85rem; }}

        .bucket-popup {{
            position: absolute; left: 0; top: 100%; z-index: 30;
            width: max-content; min-width: 100%; max-width: 500px;
            padding-top: 4px;
            display: none;  /* not just invisible: hidden popups must not widen the page */
        }}
        .bucket-popup-box {{
            background: {SURFACE_CARD}; border: 1px solid rgba(255,255,255,0.12); border-radius: 10px;
            box-shadow: 0 12px 30px rgba(0,0,0,0.5); padding: 6px;
        }}
        .bucket-popup-inner {{ max-height: 260px; overflow-y: auto; }}
        .bucket-popup-header {{
            display: flex; align-items: center; justify-content: space-between; gap: 12px;
            padding: 2px 4px 6px 10px; margin-bottom: 4px; border-bottom: 1px solid rgba(255,255,255,0.08);
            font-size: 0.72rem; font-weight: 700; letter-spacing: 0.04em; text-transform: uppercase; color: {TEXT_MUTED};
        }}
        .bucket-popup-close {{
            cursor: pointer; font-size: 1.1rem; line-height: 1; padding: 2px 8px; border-radius: 6px;
            color: {TEXT_SECONDARY}; text-transform: none;
        }}
        .bucket-popup-close:hover {{ background: rgba(255,255,255,0.1); color: {TEXT_PRIMARY}; }}
        .bucket-popup-actions {{ display: flex; align-items: center; gap: 2px; flex-shrink: 0; }}
        .bucket-popup-download {{
            cursor: pointer; display: flex; align-items: center; padding: 4px 7px; border-radius: 6px;
            color: {TEXT_SECONDARY};
        }}
        .bucket-popup-download:hover {{ background: rgba(255,255,255,0.1); color: {TEXT_PRIMARY}; }}
        .bucket-popup-box:not(:has(.bucket-popup-item)) .bucket-popup-download {{ display: none; }}
        .popup-wrap {{ position: relative; }}
        .bucket-row.has-popup {{ cursor: pointer; }}
        .bucket-row.has-popup:hover {{ background: rgba(255,255,255,0.04); }}
        .bucket-row.has-popup .label {{ cursor: pointer; }}
        .popup-wrap.open .bucket-popup {{ display: block; z-index: 40; }}
        .bucket-popup.align-right {{ left: auto; right: 0; }}
        .bucket-popup-item {{
            display: flex; align-items: center; justify-content: space-between; gap: 14px;
            padding: 6px 10px; border-radius: 6px; text-decoration: none !important;
            font-size: 0.82rem; color: {TEXT_SECONDARY};
        }}
        .bucket-popup-item:hover {{ background: rgba(255,255,255,0.06); color: {TEXT_PRIMARY}; }}
        .bucket-popup-sublink {{ padding-left: 20px; font-size: 0.74rem; opacity: 0.9; }}
        .bucket-popup-tag {{
            font-size: 0.6rem; font-weight: 700; padding: 2px 7px; border-radius: 999px;
            color: white; white-space: nowrap; flex-shrink: 0;
        }}
        .bucket-popup-key + .bucket-popup-tag {{ margin-left: -6px; margin-right: auto; }}  /* sits next to the key */
        .bucket-popup-status {{
            display: inline-flex; align-items: center; gap: 5px; flex-shrink: 0; white-space: nowrap;
            font-size: 0.66rem; font-weight: 600; color: {TEXT_SECONDARY};
            padding: 1px 8px; border-radius: 999px; border: 1px solid rgba(255,255,255,0.16);
        }}
        .bucket-popup-status .dot {{ width: 6px; height: 6px; border-radius: 50%; display: inline-block; }}
        .bucket-popup-key + .bucket-popup-status {{ margin-left: -6px; margin-right: auto; }}  /* no severity tag */
        .bucket-popup-key {{ font-weight: 700; color: {TEXT_PRIMARY}; white-space: nowrap; flex-shrink: 0; }}
        .bucket-popup-assignee {{
            text-align: right; white-space: nowrap; overflow: hidden; text-overflow: ellipsis;
            min-width: 0; max-width: 220px;
        }}
        .bucket-popup-empty {{ padding: 8px; font-size: 0.8rem; color: {TEXT_MUTED}; text-align: center; }}

        button[data-baseweb="tab"] {{
            background-color: transparent; border-radius: 999px !important; padding: 6px 20px !important;
            color: {TEXT_MUTED} !important; margin-right: 4px;
        }}
        button[data-baseweb="tab"][aria-selected="true"] {{
            background: linear-gradient(135deg, {ACCENT_FROM}, {ACCENT_TO}) !important;
            color: white !important; font-weight: 600;
        }}
        div[data-baseweb="tab-list"] {{
            gap: 2px; border-bottom: none; background: rgba(255,255,255,0.04);
            border-radius: 999px; padding: 6px; width: fit-content;
        }}
        div[data-baseweb="tab-highlight"], div[data-baseweb="tab-border"] {{ display: none; }}

        .esc-table-wrap {{
            max-height: 420px; overflow: auto; border: 1px solid rgba(255,255,255,0.07); border-radius: 10px;
        }}
        .esc-table {{ width: 100%; border-collapse: collapse; font-size: 0.82rem; color: {TEXT_SECONDARY}; }}
        .esc-table th {{
            position: sticky; top: 0; z-index: 1; background: {HEADER_BG}; color: {TEXT_MUTED};
            text-align: left; font-weight: 600; padding: 9px 12px; white-space: nowrap;
            border-bottom: 1px solid rgba(255,255,255,0.1);
        }}
        .esc-table td {{ padding: 8px 12px; border-top: 1px solid rgba(255,255,255,0.06); vertical-align: top; }}
        .esc-table tbody tr:hover {{ background: rgba(255,255,255,0.03); }}
        .esc-num {{ color: {TEXT_MUTED}; text-align: right; width: 1%; }}
        .esc-summary {{ max-width: 380px; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }}
        .esc-due {{ white-space: nowrap; }}
        .esc-table a.esc-key {{ color: #6ea8ff; font-weight: 700; text-decoration: none; white-space: nowrap; }}
        .esc-table a.esc-key:hover {{ text-decoration: underline; }}
        .esc-links {{ min-width: 110px; }}
        .esc-table a.esc-link {{
            display: inline-block; margin: 0 4px 4px 0; padding: 1px 8px; border-radius: 999px;
            background: rgba(110,168,255,0.12); color: #9cc3ff; font-size: 0.74rem; font-weight: 600;
            text-decoration: none; white-space: nowrap;
        }}
        .esc-table a.esc-link:hover {{ background: rgba(110,168,255,0.25); color: {TEXT_PRIMARY}; }}
        .esc-none {{ color: {TEXT_MUTED}; }}

        div[data-testid="stDataFrame"] {{
            border: 1px solid rgba(255,255,255,0.07); border-radius: 10px; overflow: hidden;
        }}
    </style>
    """,
    unsafe_allow_html=True,
)
st.markdown(f"<style>{ui_table.TABLE_CSS}</style>", unsafe_allow_html=True)
st.html(POPUP_SCRIPT, unsafe_allow_javascript=True)

ESCALATION_PATH_FIELD = "customfield_10883"
ESCALATION_PATH_OTHER_FIELD = "customfield_10884"


def escalation_path_value(fields):
    path_field = fields.get(ESCALATION_PATH_FIELD)
    path = path_field.get("value", "") if isinstance(path_field, dict) else (path_field or "")
    if path.strip().lower() == "other":
        other = fields.get(ESCALATION_PATH_OTHER_FIELD) or ""
        return f"Other: {other}" if other else "Other"
    return path


def extract_linked_issues(fields):
    """Pulls the linked issues (any link type/direction) off a Jira issue's
    issuelinks field into a flat, de-duplicated [{key, summary}] list."""
    links = fields.get("issuelinks") or []
    seen = set()
    result = []
    for link in links:
        linked = link.get("outwardIssue") or link.get("inwardIssue")
        if not linked:
            continue
        key = linked.get("key")
        if not key or key in seen:
            continue
        seen.add(key)
        result.append({"key": key, "summary": (linked.get("fields") or {}).get("summary", "")})
    return result


@st.cache_data(ttl=AUTO_REFRESH_SECONDS, show_spinner=False)  # a spinner per load pushes the page down -> visible shaking
def load_category_data(category, cat_config):
    fields = ["summary", "status", "assignee", "priority", "created", "issuelinks",
              ESCALATION_PATH_FIELD, ESCALATION_PATH_OTHER_FIELD]
    if cat_config["breach_type"] == "sla_from_created":
        fields.append(cat_config["severity_field"])
    elif cat_config["breach_type"] == "date_field":
        fields.append(cat_config["breach_field"])
    issues = jira_client.search_issues(cat_config["jql"], fields=fields)

    rows = []
    for issue in issues:
        f = issue["fields"]
        result = breach_logic.evaluate(issue, cat_config)
        rows.append({
            "Key": issue["key"],
            "Summary": f.get("summary", ""),
            "Status": (f.get("status") or {}).get("name", ""),
            "Priority": (f.get("priority") or {}).get("name", ""),
            "Severity": result["severity"] or "N/A",
            "Assignee": (f.get("assignee") or {}).get("displayName", "Unassigned"),
            "Created": f.get("created", ""),
            "Due": result["due_date"].strftime("%Y-%m-%d %H:%M") if result["due_date"] else "N/A",
            "Due Datetime": result["due_date"],
            "Hours Remaining": result["hours_remaining"],
            "Breach Status": result["status"],
            "Escalation Path": escalation_path_value(f),
            "Linked Issues": extract_linked_issues(f),
        })
    return pd.DataFrame(rows)


@st.cache_data(ttl=AUTO_REFRESH_SECONDS, show_spinner=False)  # a spinner per load pushes the page down -> visible shaking
def load_resolved_data(category, cat_config):
    fields = ["summary", "status", "assignee", "created", "resolutiondate", cat_config["severity_field"]]
    issues = jira_client.search_issues(cat_config["resolved_jql"], fields=fields)

    rows = []
    for issue in issues:
        f = issue["fields"]
        result = breach_logic.evaluate_resolution(issue, cat_config)
        rows.append({
            "Key": issue["key"],
            "Summary": f.get("summary", ""),
            "Status": (f.get("status") or {}).get("name", ""),
            "Severity": result["severity"] or "N/A",
            "Assignee": (f.get("assignee") or {}).get("displayName", "Unassigned"),
            "Resolved Status": result["status"],
            "Due Datetime": result["due_date"],
            "Created Today": created_today(f.get("created")),
        })
    return pd.DataFrame(rows)


@st.cache_data(ttl=AUTO_REFRESH_SECONDS, show_spinner=False)  # a spinner per load pushes the page down -> visible shaking
def load_ticket_list(jql):
    """Key/status/assignee/severity for every ticket a JQL returns (inflow lists)."""
    issues = jira_client.search_issues(jql, fields=["summary", "assignee", "status", config.SEVERITY_FIELD])
    rows = [
        {
            "Key": issue["key"],
            "Summary": issue["fields"].get("summary", ""),
            "Status": (issue["fields"].get("status") or {}).get("name", ""),
            "Assignee": (issue["fields"].get("assignee") or {}).get("displayName", "Unassigned"),
            "Severity": breach_logic._extract_option_value(issue["fields"].get(config.SEVERITY_FIELD)) or "N/A",
        }
        for issue in issues
    ]
    return pd.DataFrame(rows, columns=["Key", "Summary", "Status", "Assignee", "Severity"])


def created_today(created_raw):
    """True if the ticket was created today, using the same day boundary as
    JQL startofday(): the timezone Jira reported the timestamp in."""
    created_dt = breach_logic._parse_datetime(created_raw)
    if created_dt is None:
        return False
    return created_dt.date() == datetime.now(created_dt.tzinfo).date()


def popup_ticket(row):
    due = row.get("Due")
    if due is None and row.get("Due Datetime") is not None:
        due = breach_logic.format_due_datetime(row.get("Due Datetime"))
    return {
        "key": row["Key"],
        "assignee": row["Assignee"],
        "severity": row.get("Severity") or "N/A",
        "status": row.get("Status") or "",
        "summary": row.get("Summary") or "",
        "due": due or "",
    }


def tickets_list(df):
    return [popup_ticket(row) for _, row in df.iterrows()]


def open_split(df):
    """Splits open tickets into those created today (today's inflow) and those
    created on an earlier day (backlog)."""
    if df.empty:
        return [], []
    today = df["Created"].map(created_today)
    return tickets_list(df[today]), tickets_list(df[~today])


def split_rows(category, verb, from_inflow, from_backlog):
    """The "Today's inflow" / "Backlog" row pair shared by the open and
    resolved summary cards; each row lists its tickets on hover."""
    return (
        popup_row(
            f'<span class="label"><span class="dot" style="background:{COLOR_INFLOW};"></span>Today\'s inflow</span>'
            f'<span class="value">{len(from_inflow)}</span>',
            bucket_popup_html(from_inflow),
            f"{category} — {verb} from today's inflow",
        )
        + popup_row(
            f'<span class="label"><span class="dot" style="background:{COLOR_BACKLOG};"></span>Backlog</span>'
            f'<span class="value">{len(from_backlog)}</span>',
            bucket_popup_html(from_backlog),
            f"{category} — {verb} from backlog",
        )
    )


def mini_stat(value, label, color=TEXT_PRIMARY):
    return (
        f'<div class="mini-stat"><div class="mini-stat-value" style="color:{color};">{value}</div>'
        f'<div class="mini-stat-label">{label}</div></div>'
    )


def todays_inflow_resolved(df):
    """Resolved-today tickets that were also created today."""
    return df[df["Created Today"]] if not df.empty else df


def resolved_split(df):
    """Splits today's resolved tickets into those that were also created today
    (today's inflow) and those created on an earlier day (backlog)."""
    if df.empty:
        return [], []
    return tickets_list(df[df["Created Today"]]), tickets_list(df[~df["Created Today"]])


def resolved_tickets_by_status(df, status):
    if df.empty:
        return []
    return tickets_list(df[df["Resolved Status"] == status])


def compute_resolved_counts(df):
    if df.empty:
        return {"within_sla": 0, "breached": 0, "pct": None}
    counts = df["Resolved Status"].value_counts()
    within_sla = int(counts.get(breach_logic.STATUS_OK, 0))
    breached = int(counts.get(breach_logic.STATUS_BREACHED, 0))
    total = within_sla + breached
    pct = round(within_sla / total * 100, 1) if total else None
    return {"within_sla": within_sla, "breached": breached, "pct": pct}


def sla_pct_color(pct):
    if pct is None:
        return TEXT_MUTED
    if pct >= 90:
        return COLOR_GOOD
    if pct >= 70:
        return COLOR_WARNING
    return COLOR_CRITICAL


# ---------------------------------------------------------------------------
# Fetch all category data once (shared by every section below)
# ---------------------------------------------------------------------------
loader = st.empty()
if not st.session_state.get("dashboard_loaded"):
    loader.markdown(LOADER_HTML, unsafe_allow_html=True)

category_data = {}
category_errors = {}
resolved_data = {}
resolved_errors = {}
inflow_data = {}
inflow_errors = {}
backlog_data = {}  # only for teams with a separate backlog_jql (see config.py)
backlog_errors = {}
for category, cat_config in config.CATEGORIES.items():
    try:
        category_data[category] = load_category_data(category, cat_config)
    except Exception as exc:
        category_errors[category] = exc
    try:
        resolved_data[category] = load_resolved_data(category, cat_config)
    except Exception as exc:
        resolved_errors[category] = exc
    try:
        inflow_data[category] = load_ticket_list(cat_config["inflow_jql"])
    except Exception as exc:
        inflow_errors[category] = exc
    if "backlog_jql" in cat_config:
        try:
            # Full ticket rows (severity, due, breach status), so the backlog
            # can be bucketed in Time to Resolve like the team's other tickets.
            backlog_data[category] = load_category_data(category, {**cat_config, "jql": cat_config["backlog_jql"]})
        except Exception as exc:
            backlog_errors[category] = exc

loader.empty()
st.session_state["dashboard_loaded"] = True

for category, exc in category_errors.items():
    st.error(f"Failed to load {category}: {exc}")
for category, exc in resolved_errors.items():
    st.error(f"Failed to load resolved tickets for {category}: {exc}")
for category, exc in inflow_errors.items():
    st.error(f"Failed to load today's inflow for {category}: {exc}")
for category, exc in backlog_errors.items():
    st.error(f"Failed to load backlog for {category}: {exc}")

ok_categories = {c: df for c, df in category_data.items() if c not in category_errors}

# ---------------------------------------------------------------------------
# Sidebar: chat assistant (Claude, tool use over the data already loaded above)
# Disabled for now (no ANTHROPIC_API_KEY set up yet) -- re-enable by
# uncommenting once Anthropic Console billing + a key are in place.
# ---------------------------------------------------------------------------
# with st.sidebar:
#     st.markdown("### Ask about your tickets")
#     if not config.ANTHROPIC_API_KEY:
#         st.caption("Set ANTHROPIC_API_KEY in config.py to enable this.")
#     else:
#         if "ai_display" not in st.session_state:
#             st.session_state.ai_display = []
#         if "ai_history" not in st.session_state:
#             st.session_state.ai_history = []
#
#         for turn in st.session_state.ai_display:
#             with st.chat_message(turn["role"]):
#                 st.markdown(turn["text"])
#
#         question = st.chat_input("e.g. what's breaching in B2B right now?")
#         if question:
#             with st.chat_message("user"):
#                 st.markdown(question)
#             st.session_state.ai_display.append({"role": "user", "text": question})
#
#             with st.chat_message("assistant"):
#                 with st.spinner("Thinking..."):
#                     try:
#                         answer, updated_history = ai_assistant.ask(
#                             question, st.session_state.ai_history, category_data, resolved_data
#                         )
#                     except Exception as exc:
#                         answer = f"Sorry, something went wrong: {exc}"
#                         updated_history = st.session_state.ai_history
#                 st.markdown(answer)
#             st.session_state.ai_history = updated_history
#             st.session_state.ai_display.append({"role": "assistant", "text": answer})

# ---------------------------------------------------------------------------
# Header: title, overall resolved/SLA summary, refresh button
# ---------------------------------------------------------------------------
# Open tickets bucketed by time to resolve, per team. A team whose open query
# only covers today's tickets also gets its backlog (older open tickets), so
# past breaches show up too. Used by Time to Resolve and by Today SLA.
open_buckets = {}
for category in config.CATEGORIES:
    if category in category_errors:
        continue
    open_df = category_data[category]
    if category in backlog_data:
        open_df = pd.concat([open_df, backlog_data[category]], ignore_index=True)
    open_buckets[category] = compute_time_buckets(open_df)


def today_sla_pct(within_sla, breached_resolved, breached_open):
    """Today SLA: tickets resolved today within SLA, out of everything judged
    today -- resolved today (within or breached) plus open tickets whose SLA
    ran out today."""
    judged = within_sla + breached_resolved + breached_open
    return round(within_sla / judged * 100, 1) if judged else None


# Today SLA is judged on tickets created and resolved today, plus open tickets
# that breached today; tickets resolved from past inflow are counted separately.
today_within_sla = 0
today_breached = 0
today_breached_open = 0
resolved_today_inflow_count = 0
resolved_past_inflow_count = 0
for category, df in resolved_data.items():
    if category in resolved_errors:
        continue
    today_df = todays_inflow_resolved(df)
    counts = compute_resolved_counts(today_df)
    today_within_sla += counts["within_sla"]
    today_breached += counts["breached"]
    resolved_today_inflow_count += len(today_df)
    resolved_past_inflow_count += len(df) - len(today_df)
    if category in open_buckets:
        today_breached_open += open_buckets[category][BUCKET_BREACHED_TODAY]["count"]
today_pct = today_sla_pct(today_within_sla, today_breached, today_breached_open)
today_pct_display = "—" if today_pct is None else f"{today_pct:.0f}%"

# Team scopes can overlap (e.g. SAC covers every non-Retail SAC ticket), so
# the overall inflow counts each ticket once rather than summing the teams.
overall_inflow_keys = set()
for category, df in inflow_data.items():
    overall_inflow_keys.update(df["Key"])
overall_inflow_display = "—" if inflow_errors and not inflow_data else len(overall_inflow_keys)

header_col1, header_col2, header_col3 = st.columns([2.5, 3.9, 1.3])
with header_col1:
    st.markdown(
        HEADER_HTML,
        unsafe_allow_html=True,
    )
with header_col2:
    st.markdown(
        '<div class="overall-summary">'
        '<div class="overall-stat" title="Tickets created today across all teams (each ticket counted once)">'
        f'<div class="overall-stat-value">{overall_inflow_display}</div>'
        '<div class="overall-stat-label">Inflow Today</div></div>'
        '<div class="overall-divider"></div>'
        '<div class="overall-stat" title="Tickets created before today and resolved today">'
        f'<div class="overall-stat-value">{resolved_past_inflow_count}</div>'
        '<div class="overall-stat-label">Resolved from Past Inflow</div></div>'
        '<div class="overall-divider"></div>'
        '<div class="overall-stat" title="Tickets created today and already resolved">'
        f'<div class="overall-stat-value">{resolved_today_inflow_count}</div>'
        '<div class="overall-stat-label">Resolved from Today Inflow</div></div>'
        '<div class="overall-divider"></div>'
        '<div class="overall-stat" title="Within SLA / (Within SLA + Resolved late + Open, breached today)">'
        f'<div class="overall-stat-value" style="color:{sla_pct_color(today_pct)};">{today_pct_display}</div>'
        '<div class="overall-stat-label">Today SLA</div></div>'
        '</div>',
        unsafe_allow_html=True,
    )
with header_col3:
    if st.button("🔄  Refresh now", use_container_width=True):
        st.cache_data.clear()
        st.rerun()
    # pages/Ticket_Trend_Analysis.py and pages/Resolved_Tickets_History.py,
    # each opened in its own browser tab.
    st.markdown(
        '<div class="page-links">'
        '<a class="trend-link" href="Ticket_Trend_Analysis" target="_blank" rel="noopener">'
        '📈&nbsp; Ticket Trend Analysis &#8599;</a>'
        '<a class="trend-link" href="Resolved_Tickets_History" target="_blank" rel="noopener">'
        '✅&nbsp; Resolved History &#8599;</a></div>'
        # Same block as the links, so the note always sits below them.
        f'<div class="app-refreshed">Last refreshed: {breach_logic.jira_now():%Y-%m-%d %H:%M:%S %Z}'
        f' &middot; auto-refreshes every {AUTO_REFRESH_SECONDS // 60} min'
        f'<br>All times in {html.escape(breach_logic.jira_time_label())}</div>',
        unsafe_allow_html=True,
    )

st.markdown('<div class="header-spacer"></div>', unsafe_allow_html=True)

# ---------------------------------------------------------------------------
# Section: Total Open Tickets  |  Total Resolved Tickets - Today
# ---------------------------------------------------------------------------
open_section, resolved_section = st.columns(2, gap="large")

with open_section:
    st.markdown('<div class="section-title">Total Open Tickets &amp; Inflow &mdash; Today</div>', unsafe_allow_html=True)
    cols = st.columns(len(config.CATEGORIES))
    for col, category in zip(cols, config.CATEGORIES.keys()):
        if category in inflow_errors:
            inflow_stat = mini_stat("—", "Inflow today")
        else:
            inflow_stat = popup_row(
                f'<div class="mini-stat-value">{len(inflow_data[category])}</div>'
                f'<div class="mini-stat-label">Inflow today</div>',
                bucket_popup_html(tickets_list(inflow_data[category])),
                f"{category} — Inflow today",
                row_class="mini-stat",
            )
        if category in category_errors or category in backlog_errors:
            open_value = "—"
            split_html = '<div class="bucket-row"><span class="label">Error loading</span></div>'
        else:
            open_inflow, open_backlog = open_split(category_data[category])
            if category in backlog_data:
                # This team's open query only covers today's tickets; its
                # backlog comes from backlog_jql instead.
                open_backlog = tickets_list(backlog_data[category])
            open_value = len(open_inflow) + len(open_backlog)
            split_html = (
                '<div class="bucket-subcaption">Open from</div>'
                + split_rows(category, "Open", open_inflow, open_backlog)
            )
        col.markdown(
            f'<div class="bucket-card summary-card"><div class="bucket-card-head">{category}</div>'
            f'<div class="mini-stats mini-stats-grow">{mini_stat(open_value, "Open now")}'
            f'<div class="mini-stat-divider"></div>{inflow_stat}</div>'
            f'{split_html}</div>',
            unsafe_allow_html=True,
        )

with resolved_section:
    st.markdown(
        section_title_with_download(
            "Total Resolved Tickets &mdash; Today's Inflow", RESOLVED_BREACHED_GROUP,
            "Download breached", "Breached_todays_inflow_all_teams",
        ),
        unsafe_allow_html=True,
    )
    cols = st.columns(len(config.CATEGORIES))
    for col, category in zip(cols, config.CATEGORIES.keys()):
        if category in resolved_errors:
            col.markdown(
                f'<div class="bucket-card summary-card"><div class="bucket-card-head">{category}</div>'
                f'<div class="bucket-row"><span class="label">Error loading</span></div></div>',
                unsafe_allow_html=True,
            )
            continue

        # The card is about today's inflow: tickets created today and resolved
        # today. Past-inflow resolutions are shown once, at the bottom.
        from_inflow, from_backlog = resolved_split(resolved_data[category])
        today_df = todays_inflow_resolved(resolved_data[category])
        counts = compute_resolved_counts(today_df)
        breached_open = open_buckets.get(category, {}).get(BUCKET_BREACHED_TODAY, {"count": 0, "tickets": []})
        pct = today_sla_pct(counts["within_sla"], counts["breached"], breached_open["count"])
        pct_display = "—" if pct is None else f"{pct:.0f}%"
        rows_html = (
            '<div class="mini-stats">'
            + popup_row(
                f'<div class="mini-stat-value">{len(from_inflow)}</div>'
                f'<div class="mini-stat-label">Resolved</div>',
                bucket_popup_html(from_inflow),
                f"{category} — Resolved from today's inflow",
                row_class="mini-stat",
            )
            + '<div class="mini-stat-divider"></div>'
            + mini_stat(pct_display, "Today SLA", sla_pct_color(pct))
            + '</div>'
            + popup_row(
                f'<span class="label"><span class="dot" style="background:{COLOR_GOOD};"></span>Within SLA</span>'
                f'<span class="value">{counts["within_sla"]}</span>',
                bucket_popup_html(resolved_tickets_by_status(today_df, breach_logic.STATUS_OK)),
                f"{category} — Within SLA (today's inflow)",
            )
            + popup_row(
                f'<span class="label"><span class="dot" style="background:{COLOR_CRITICAL};"></span>Resolved late</span>'
                f'<span class="value">{counts["breached"]}</span>',
                bucket_popup_html(resolved_tickets_by_status(today_df, breach_logic.STATUS_BREACHED)),
                f"{category} — Resolved late (today's inflow)",
                export_group=RESOLVED_BREACHED_GROUP,
            )
            + popup_row(
                f'<span class="label"><span class="dot" style="background:{BUCKET_COLOR[BUCKET_BREACHED_TODAY]};"></span>Open, breached</span>'
                f'<span class="value">{breached_open["count"]}</span>',
                bucket_popup_html(breached_open["tickets"]),
                f"{category} — Open, breached today",
                export_group=RESOLVED_BREACHED_GROUP,
            )
            + '<div class="bucket-subcaption">Also resolved today</div>'
            + popup_row(
                f'<span class="label"><span class="dot" style="background:{COLOR_BACKLOG};"></span>Past inflow</span>'
                f'<span class="value">{len(from_backlog)}</span>',
                bucket_popup_html(from_backlog),
                f"{category} — Resolved from past inflow",
            )
        )
        col.markdown(
            f'<div class="bucket-card summary-card"><div class="bucket-card-head">{category}</div>{rows_html}</div>',
            unsafe_allow_html=True,
        )

# ---------------------------------------------------------------------------
# Section: Time to Resolve - Open Tickets
# ---------------------------------------------------------------------------
st.markdown(
    section_title_with_download(
        "Time to Resolve &mdash; Open Tickets", TTR_BREACHED_GROUP,
        "Download breached", "Breached_open_tickets_all_teams",
    ),
    unsafe_allow_html=True,
)
cols = st.columns(len(config.CATEGORIES))
for col, category in zip(cols, config.CATEGORIES.keys()):
    if category in category_errors:
        col.markdown(
            f'<div class="bucket-card"><div class="bucket-card-head">{category}</div>'
            f'<div class="bucket-row"><span class="label">Error loading</span></div></div>',
            unsafe_allow_html=True,
        )
        continue

    buckets = open_buckets[category]
    rows_html = "".join(
        popup_row(
            f'<span class="label"><span class="dot" style="background:{BUCKET_COLOR[b]};"></span>{b}</span>'
            f'<span class="value">{buckets[b]["count"]}</span>',
            bucket_popup_html(buckets[b]["tickets"]),
            f"{category} — {b}",
            export_group=TTR_BREACHED_GROUP if b in (BUCKET_BREACHED_PAST, BUCKET_BREACHED_TODAY) else None,
        )
        for b in BUCKET_ORDER
    )
    col.markdown(
        f'<div class="bucket-card"><div class="bucket-card-head">{category}</div>{rows_html}</div>',
        unsafe_allow_html=True,
    )

# ---------------------------------------------------------------------------
# Section: Tickets by Status
# ---------------------------------------------------------------------------
st.markdown('<div class="section-title">Total Open Tickets by Status</div>', unsafe_allow_html=True)
cols = st.columns(len(config.CATEGORIES))
for col, category in zip(cols, config.CATEGORIES.keys()):
    if category in category_errors:
        col.markdown(
            f'<div class="bucket-card"><div class="bucket-card-head">{category}</div>'
            f'<div class="bucket-row"><span class="label">Error loading</span></div></div>',
            unsafe_allow_html=True,
        )
        continue

    status_counts = compute_status_counts(category_data[category])
    if not status_counts:
        rows_html = '<div class="bucket-row"><span class="label">No open tickets</span></div>'
    else:
        escalated_tickets = escalated_tickets_with_links(category_data[category])
        rows_html = "".join(
            popup_row(
                f'<span class="label"><span class="dot" style="background:{status_dot_color(status)};"></span>{status}</span>'
                f'<span class="value">{count}</span>',
                escalated_links_popup_html(escalated_tickets),
                f"{category} — Escalated",
            )
            if status.strip().lower() == "escalated"
            else popup_row(
                f'<span class="label"><span class="dot" style="background:{status_dot_color(status)};"></span>{status}</span>'
                f'<span class="value">{count}</span>',
                bucket_popup_html(open_tickets_with_status(category_data[category], "waiting for support")),
                f"{category} — {status}",
            )
            if status.strip().lower() == "waiting for support"
            else (
                f'<div class="bucket-row"><span class="label">'
                f'<span class="dot" style="background:{status_dot_color(status)};"></span>{status}</span>'
                f'<span class="value">{count}</span></div>'
            )
            for status, count in status_counts.items()
        )
    col.markdown(
        f'<div class="bucket-card"><div class="bucket-card-head">{category}</div>{rows_html}</div>',
        unsafe_allow_html=True,
    )

# ---------------------------------------------------------------------------
# Section: Escalated Tickets (path + severity)
# ---------------------------------------------------------------------------
st.markdown('<div class="section-title">Escalated Tickets</div>', unsafe_allow_html=True)
escalated_frames = []
for category, df in ok_categories.items():
    if df.empty:
        continue
    escalated = df[df["Status"].str.strip().str.lower() == "escalated"].copy()
    if not escalated.empty:
        escalated.insert(0, "Category", category)
        escalated_frames.append(escalated)

if escalated_frames:
    escalated_df = pd.concat(escalated_frames, ignore_index=True)
    st.markdown(escalated_table_html(escalated_df), unsafe_allow_html=True)
else:
    st.info("No tickets are currently escalated across any queue.")

# ---------------------------------------------------------------------------
# Section: Ticket Details (per-category, filterable)
# ---------------------------------------------------------------------------
st.markdown('<div class="section-title">Ticket Details</div>', unsafe_allow_html=True)

tabs = st.tabs(list(config.CATEGORIES.keys()))
for tab, (category, cat_config) in zip(tabs, config.CATEGORIES.items()):
    with tab:
        if category in category_errors:
            st.error(f"Failed to load {category}: {category_errors[category]}")
            continue
        df = category_data[category]

        if df.empty:
            st.info("No tickets matched this query.")
            continue

        breach_filter = st.multiselect(
            "Filter by breach status",
            options=STATUS_ORDER,
            default=[breach_logic.STATUS_WARNING, breach_logic.STATUS_BREACHED],
            key=f"filter_{category}",
        )
        shown = df[df["Breach Status"].isin(breach_filter)] if breach_filter else df

        st.markdown(ui_table.html_table(ticket_details_view(shown), max_height=500), unsafe_allow_html=True)

# ---------------------------------------------------------------------------
# Section: Productivity by Assignee — Today
# ---------------------------------------------------------------------------
st.markdown('<div class="section-title">Productivity by Assignee &mdash; Today</div>', unsafe_allow_html=True)

category_names = list(config.CATEGORIES.keys())
team_filter = st.multiselect(
    "Filter by team",
    options=category_names,
    default=category_names,
    key="productivity_team_filter",
)
selected_categories = [cat for cat in category_names if cat in team_filter]

per_assignee_open = defaultdict(lambda: defaultdict(int))
for category, df in category_data.items():
    if category not in selected_categories or category in category_errors or df.empty:
        continue
    for assignee, count in df["Assignee"].value_counts().items():
        per_assignee_open[assignee][category] = int(count)

per_assignee_resolved = defaultdict(lambda: defaultdict(int))
for category, df in resolved_data.items():
    if category not in selected_categories or category in resolved_errors or df.empty:
        continue
    for assignee, count in df["Assignee"].value_counts().items():
        per_assignee_resolved[assignee][category] = int(count)

assignees = set(per_assignee_open) | set(per_assignee_resolved)

if not selected_categories:
    st.info("Select at least one team to see productivity data.")
elif not assignees:
    st.info("No open or resolved tickets to report on.")
else:
    open_cols = [f"{cat} (Open)" for cat in selected_categories]
    resolved_cols = [f"{cat} (Resolved)" for cat in selected_categories]

    rows = []
    for assignee in assignees:
        row = {"Assignee": assignee}
        total_open = 0
        total_resolved = 0
        for cat in selected_categories:
            row[f"{cat} (Open)"] = per_assignee_open[assignee].get(cat, 0)
            row[f"{cat} (Resolved)"] = per_assignee_resolved[assignee].get(cat, 0)
            total_open += row[f"{cat} (Open)"]
            total_resolved += row[f"{cat} (Resolved)"]
        row["Total Open"] = total_open
        row["Total Resolved Today"] = total_resolved
        rows.append(row)

    columns = ["Assignee", *open_cols, *resolved_cols, "Total Open", "Total Resolved Today"]
    productivity_df = pd.DataFrame(rows, columns=columns)
    productivity_df = productivity_df.sort_values("Total Open", ascending=True).reset_index(drop=True)

    grand_total = {"Assignee": "Grand Total"}
    for col in [*open_cols, *resolved_cols, "Total Open", "Total Resolved Today"]:
        grand_total[col] = int(productivity_df[col].sum())
    display_df = pd.concat([productivity_df, pd.DataFrame([grand_total])], ignore_index=True)

    st.markdown(ui_table.html_table(display_df, total_row=True, max_height=500), unsafe_allow_html=True)

    excel_buffer = io.BytesIO()
    with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
        display_df.to_excel(writer, sheet_name="Productivity", index=False)
        ws = writer.sheets["Productivity"]
        for col_idx, col_name in enumerate(display_df.columns, start=1):
            max_len = max(len(str(col_name)), display_df[col_name].astype(str).map(len).max()) + 4
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max_len
        bold = Font(bold=True)
        for cell in ws[1]:
            cell.font = bold
        for cell in ws[ws.max_row]:
            cell.font = bold

    st.download_button(
        "⬇️ Download as Excel",
        data=excel_buffer.getvalue(),
        file_name=f"Productivity_Report_{datetime.now().strftime('%Y-%m-%d')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
